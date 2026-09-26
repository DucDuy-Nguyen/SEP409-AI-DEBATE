"""
Kiem tra AI Evaluator co cham ON DINH giua tieng Anh va tieng Viet hay khong —
so sanh AI VOI CHINH NO qua 2 ngon ngu input (KHONG phai so voi nguoi, khong
can ground truth). Dung lai du lieu argument tieng Anh da co san trong
data/dagstuhl_processed.csv (tu lan xu ly Dagstuhl truoc), dich sang tieng
Viet bang LLM, roi cham CA HAI ban qua evaluate_argument() de so sanh.

Cach dung:
    python scripts/run_language_consistency_check.py --n 8 --runs 2 --seed 42

    # Kiem chung sau 1 argument cu the (bo qua --n/--seed), vi du runs cao hon
    # de do do on dinh rieng cua case do:
    python scripts/run_language_consistency_check.py --focus-id 46238 --runs 5

Yeu cau: da dien LLM key that trong .env, va da co san
data/dagstuhl_processed.csv (chay scripts/run_dagstuhl_agreement_eval.py it
nhat 1 lan truoc de tao file cache nay).

TACH "LECH DO NGON NGU" KHOI "NHIEU NGAU NHIEN" (2026-09): moi ngon ngu duoc
cham --runs lan doc lap (giong co che --runs cua run_ibm_agreement_eval.py /
run_dagstuhl_agreement_eval.py), tinh duoc range (= max - min) rieng cho tung
ngon ngu — do CHINH nhieu noi tai (temperature) cua ngon ngu do. Lech giua 2
ngon ngu (diff = ai_score_avg_VI - ai_score_avg_EN) chi duoc coi la "dang ke"
(that su do ngon ngu) neu |diff| > range_EN + range_VI cong lai — tuc lon hon
tong nhieu noi tai cua CA 2 phia. Neu khong, lech quan sat duoc rat co the
chi la dao dong ngau nhien binh thuong, khong phai do khac ngon ngu.

Moi argument ton (1 + 2*runs) lenh goi LLM: 1 dich + `runs` lan cham EN +
`runs` lan cham VI (vi du runs=2 -> 5 lenh/argument).

QUAN TRONG: giong cac script truoc, neu bi loi/ngat giua chung, script VAN
XUAT ra file Excel voi nhung argument da xu ly xong, khong mat trang.
"""
import argparse
import json
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.config import get_settings, Settings
from app.schemas import ArgumentEvaluationRequest, DebateSide, DebateStage
from app.services.evaluator import evaluate_argument
from app.services.json_utils import EvaluationParseError, call_llm_with_retry
from app.services.llm_client import LLMClient, get_llm_client

PROCESSED_CSV_PATH = Path(__file__).parent.parent / "data" / "dagstuhl_processed.csv"
OUTPUT_PATH = Path(__file__).parent.parent / "language_consistency_results.xlsx"
DELAY_BETWEEN_CALLS_SEC = 1.5

AI_FIELDS = ["overall_score", "Logic", "Evidence", "Relevance", "Structure", "Persuasiveness"]

# Nguong "range_EN + range_VI == 0" coi la khong du du lieu (thay vi TRUE/FALSE) —
# dung epsilon nho de tranh sai so lam tron float thay vi so sanh == 0 tuyet doi.
_ZERO_RANGE_EPSILON = 1e-9


TRANSLATE_SYSTEM_PROMPT = """Ban la 1 bien dich vien chuyen nghiep, dich noi dung tranh bien tu tieng Anh \
sang tieng Viet.

NHIEM VU: Dich CHINH XAC va tu nhien, GIU NGUYEN y nghia, sac thai, va muc do \
manh/yeu cua lap luan trong ban goc — KHONG dien giai them, KHONG tom tat, \
KHONG bo sot y, KHONG them binh luan hay sua loi lap luan cua ban goc (ke ca \
neu ban goc co lap luan yeu hoac sai — dich dung y, khong "sua ho").

Dich CA HAI: "motion" (chu de tranh bien) va "argument" (noi dung lap luan).

[OUTPUT - CHI tra ve JSON hop le, khong them text nao khac ngoai JSON]
{
  "motion_vi": "<ban dich tieng Viet cua motion>",
  "argument_vi": "<ban dich tieng Viet cua argument, giu nguyen so doan/cau truc neu co>"
}"""


def build_translate_user_prompt(motion: str, argument_text: str) -> str:
    return f"[MOTION]\n{motion}\n\n[ARGUMENT]\n{argument_text}"


def _build_translation_result(parsed: dict, raw_text: str) -> tuple[str, str]:
    motion_vi = parsed.get("motion_vi")
    argument_vi = parsed.get("argument_vi")
    if not motion_vi or not argument_vi:
        raise EvaluationParseError(
            "LLM khong tra ve du 'motion_vi' va 'argument_vi' khi dich — coi la loi dinh dang de retry."
        )
    return motion_vi, argument_vi


def translate_to_vietnamese(motion: str, argument_text: str, client: LLMClient) -> tuple[str, str]:
    """Dich (motion, argument_text) sang tieng Viet, tai su dung call_llm_with_retry
    (json_utils.py) de retry neu LLM tra JSON khong hop le, giong cac evaluator khac."""
    user_prompt = build_translate_user_prompt(motion, argument_text)
    return call_llm_with_retry(
        generate_fn=lambda up: client.generate(TRANSLATE_SYSTEM_PROMPT, up),
        system_prompt=TRANSLATE_SYSTEM_PROMPT,
        user_prompt=user_prompt,
        build_result_fn=_build_translation_result,
    )


def load_dagstuhl_sample(n: int, seed: int = 42) -> pd.DataFrame:
    if not PROCESSED_CSV_PATH.exists():
        raise FileNotFoundError(
            f"Khong tim thay {PROCESSED_CSV_PATH}. Chay scripts/run_dagstuhl_agreement_eval.py "
            "it nhat 1 lan truoc de tao file cache nay."
        )
    df = pd.read_csv(PROCESSED_CSV_PATH)
    n = min(n, len(df))
    return df.sample(n=n, random_state=seed).reset_index(drop=True)


def _normalize_dagstuhl_id(raw_id: str) -> str:
    """Dataset co CA 2 dang id: 'arg219250' va id so tran nhu '46238' (da phat
    hien khi trinh sat du lieu Dagstuhl). Bo tien to 'arg' (neu co) de so khop
    linh hoat giua 2 dang."""
    raw_id = str(raw_id).strip()
    return raw_id[3:] if raw_id.lower().startswith("arg") else raw_id


def load_dagstuhl_single(focus_id: str) -> pd.DataFrame:
    """
    Lay DUNG 1 dong co #id khop `focus_id` — uu tien khop CHINH XAC chuoi id
    truoc, neu khong co moi thu khop SAU KHI chuan hoa (bo tien to 'arg' o CA
    HAI phia) de xu ly duoc ca 2 dang id trong dataset. Dung cho --focus-id
    (kiem chung sau 1 argument cu the, tach biet voi batch ngau nhien).
    """
    if not PROCESSED_CSV_PATH.exists():
        raise FileNotFoundError(
            f"Khong tim thay {PROCESSED_CSV_PATH}. Chay scripts/run_dagstuhl_agreement_eval.py "
            "it nhat 1 lan truoc de tao file cache nay."
        )
    df = pd.read_csv(PROCESSED_CSV_PATH)
    df["id"] = df["id"].astype(str)

    exact = df[df["id"] == focus_id]
    if len(exact) == 1:
        return exact.reset_index(drop=True)

    target_normalized = _normalize_dagstuhl_id(focus_id)
    matched = df[df["id"].apply(_normalize_dagstuhl_id) == target_normalized]

    if matched.empty:
        raise ValueError(
            f"Khong tim thay argument nao voi #id khop '{focus_id}' "
            f"(da thu ca dang co/khong tien to 'arg') trong {PROCESSED_CSV_PATH}."
        )
    if len(matched) > 1:
        print(
            f"CANH BAO: co {len(matched)} dong khop voi id '{focus_id}' sau khi chuan hoa "
            f"tien to 'arg' — dung dong DAU TIEN: id={matched.iloc[0]['id']}"
        )
    return matched.iloc[[0]].reset_index(drop=True)


def _run_language_n_times(
    motion: str, argument_text: str, settings: Settings, client: LLMClient, n_runs: int
) -> tuple[dict[str, list[float]], list[str]]:
    """Cham 1 ban (1 ngon ngu) cua 1 argument `n_runs` lan doc lap. Tra ve
    (scores_by_field, reasoning_list) — scores_by_field[field] la list cac
    diem thu duoc qua cac lan chay THANH CONG (co the < n_runs neu co lan loi)."""
    req = ArgumentEvaluationRequest(
        motion=motion,
        # Context SOLO (opponent_argument_text=None) -> side khong anh huong
        # ban chat cham diem, xem giai thich trong cac script truoc.
        side=DebateSide.PRO,
        stage=DebateStage.OPENING,
        argument_text=argument_text,
        opponent_argument_text=None,
    )
    scores_by_field: dict[str, list[float]] = {f: [] for f in AI_FIELDS}
    reasoning_list: list[str] = []

    for run_idx in range(1, n_runs + 1):
        try:
            result = evaluate_argument(req, settings=settings, llm_client=client)
            scores_by_field["overall_score"].append(result.overall_score)
            for c in result.criteria:
                if c.name in scores_by_field:
                    scores_by_field[c.name].append(c.score)
            reasoning_list.append(result.overall_reasoning)
            print(f"run{run_idx}=OK", end=" ")
        except Exception as e:
            print(f"run{run_idx}=LOI({type(e).__name__})", end=" ")
        if run_idx < n_runs:
            time.sleep(DELAY_BETWEEN_CALLS_SEC)

    return scores_by_field, reasoning_list


def _aggregate_avg_range(scores_by_field: dict[str, list[float]]) -> tuple[dict[str, float | None], dict[str, float | None]]:
    avg: dict[str, float | None] = {}
    rng: dict[str, float | None] = {}
    for field, values in scores_by_field.items():
        if values:
            avg[field] = round(sum(values) / len(values), 2)
            rng[field] = round(max(values) - min(values), 2)
        else:
            avg[field] = None
            rng[field] = None
    return avg, rng


def _compute_diff_and_significance(
    en_avg: dict[str, float | None], en_range: dict[str, float | None],
    vi_avg: dict[str, float | None], vi_range: dict[str, float | None],
) -> tuple[dict[str, float | None], dict[str, object]]:
    """
    diff[field] = vi_avg - en_avg (None neu thieu 1 trong 2 phia).
    significant[field]:
      - "Khong du du lieu" neu thieu diff, hoac range_EN + range_VI xap xi 0
        (vi du chi co 1 lan chay thanh cong o 1 trong 2 phia -> khong do duoc
        nhieu noi tai de lam moc so sanh).
      - True/False: |diff| > (range_EN + range_VI) hay khong.
    """
    diff: dict[str, float | None] = {}
    significant: dict[str, object] = {}

    for field in AI_FIELDS:
        e_avg, v_avg = en_avg.get(field), vi_avg.get(field)
        e_rng, v_rng = en_range.get(field), vi_range.get(field)

        if e_avg is None or v_avg is None:
            diff[field] = None
            significant[field] = "Khong du du lieu"
            continue

        diff[field] = round(v_avg - e_avg, 2)

        if e_rng is None or v_rng is None:
            significant[field] = "Khong du du lieu"
            continue

        total_range = e_rng + v_rng
        if total_range < _ZERO_RANGE_EPSILON:
            significant[field] = "Khong du du lieu"
        else:
            significant[field] = abs(diff[field]) > total_range

    return diff, significant


def _evaluate_one_row(row, settings: Settings, client: LLMClient, n_runs: int) -> dict:
    """
    (1 + 2*n_runs) lenh goi LLM cho 1 argument: dich -> cham EN n_runs lan ->
    cham VI n_runs lan. Moi buoc lon co try/except RIENG — neu dich loi thi
    khong cham duoc ca 2 ban (tra ve som kem error), neu ca n_runs lan cham EN
    deu that bai thi khong co gi de so sanh (tra ve som).
    """
    entry = {
        "id": row["id"],
        "motion_en": row["issue"],
        "text_en": row["text"],
        "motion_vi": None,
        "argument_vi": None,
        "en_avg": None, "en_range": None,
        "vi_avg": None, "vi_range": None,
        "en_overall_reasoning": None,
        "vi_overall_reasoning": None,
        "diff": None,
        "significant": None,
        "error": None,
    }

    try:
        motion_vi, argument_vi = translate_to_vietnamese(entry["motion_en"], entry["text_en"], client)
        entry["motion_vi"] = motion_vi
        entry["argument_vi"] = argument_vi
    except Exception as e:
        entry["error"] = f"Dich loi: {e}"
        return entry
    time.sleep(DELAY_BETWEEN_CALLS_SEC)

    print("[EN]", end=" ")
    en_scores, en_reasoning_list = _run_language_n_times(
        entry["motion_en"], entry["text_en"], settings, client, n_runs
    )
    en_avg, en_range = _aggregate_avg_range(en_scores)
    if en_avg["overall_score"] is None:
        entry["error"] = "Cham tieng Anh that bai toan bo cac lan chay"
        return entry
    entry["en_avg"], entry["en_range"] = en_avg, en_range
    entry["en_overall_reasoning"] = en_reasoning_list[0] if en_reasoning_list else None
    time.sleep(DELAY_BETWEEN_CALLS_SEC)

    print("| [VI]", end=" ")
    vi_scores, vi_reasoning_list = _run_language_n_times(
        entry["motion_vi"], entry["argument_vi"], settings, client, n_runs
    )
    vi_avg, vi_range = _aggregate_avg_range(vi_scores)
    if vi_avg["overall_score"] is None:
        entry["error"] = "Cham tieng Viet that bai toan bo cac lan chay"
        return entry
    entry["vi_avg"], entry["vi_range"] = vi_avg, vi_range
    entry["vi_overall_reasoning"] = vi_reasoning_list[0] if vi_reasoning_list else None

    entry["diff"], entry["significant"] = _compute_diff_and_significance(en_avg, en_range, vi_avg, vi_range)
    return entry


def run_batch(df: pd.DataFrame, settings: Settings, client: LLMClient, n_runs: int, rows: list) -> None:
    """Ghi truc tiep vao `rows` truyen vao — giong cac script truoc, de neu
    loi/Ctrl+C giua chung, ket qua da xu ly duoc khong bi mat."""
    for i, row in df.iterrows():
        print(f"[{i + 1}/{len(df)}] id={row['id']} ...", end=" ")
        try:
            entry = _evaluate_one_row(row, settings, client, n_runs)
            rows.append(entry)
            if entry["error"]:
                print(f"\n    LOI: {entry['error']}")
            else:
                d = entry["diff"]["overall_score"]
                sig = entry["significant"]["overall_score"]
                print(
                    f"\n    overall: EN_avg={entry['en_avg']['overall_score']} "
                    f"(range={entry['en_range']['overall_score']}) "
                    f"VI_avg={entry['vi_avg']['overall_score']} "
                    f"(range={entry['vi_range']['overall_score']}) "
                    f"diff={d:+.2f} dang_ke={sig}"
                )
        except Exception as e:
            print(f"LOI TOAN BO DONG: {e}")
            rows.append({
                "id": row.get("id"), "motion_en": row.get("issue"), "text_en": row.get("text"),
                "motion_vi": None, "argument_vi": None,
                "en_avg": None, "en_range": None, "vi_avg": None, "vi_range": None,
                "en_overall_reasoning": None, "vi_overall_reasoning": None,
                "diff": None, "significant": None, "error": str(e),
            })


def compute_summary_metrics(rows: list[dict]) -> dict[str, dict]:
    """
    Cho MOI field trong AI_FIELDS: mean diff co dau, MAE, range_EN/range_VI
    trung binh, va % argument co "Lech co dang ke" = True (tren tong so
    argument hop le co ca diff lan range o CA 2 ngon ngu). Pearson r CHI tinh
    cho overall_score (thang 0-10 lien tuc, y nghia hon thang 1-5 hep o n nho).
    """
    from scipy import stats

    valid = [r for r in rows if r.get("en_avg") is not None and r.get("vi_avg") is not None]
    summary = {}

    for field in AI_FIELDS:
        n = len(valid)
        if n == 0:
            summary[field] = {
                "n": 0, "mean_diff_vi_minus_en": None, "mae": None,
                "pearson_r": None, "p_value": None,
                "mean_range_en": None, "mean_range_vi": None, "pct_significant": None,
            }
            continue

        diffs = [r["diff"][field] for r in valid if r["diff"].get(field) is not None]
        mean_diff = round(sum(diffs) / len(diffs), 2) if diffs else None
        mae = round(sum(abs(d) for d in diffs) / len(diffs), 2) if diffs else None

        en_ranges = [r["en_range"][field] for r in valid if r["en_range"].get(field) is not None]
        vi_ranges = [r["vi_range"][field] for r in valid if r["vi_range"].get(field) is not None]
        mean_range_en = round(sum(en_ranges) / len(en_ranges), 2) if en_ranges else None
        mean_range_vi = round(sum(vi_ranges) / len(vi_ranges), 2) if vi_ranges else None

        sig_flags = [r["significant"][field] for r in valid if field in r.get("significant", {})]
        # Mau so = tat ca argument hop le (bao gom ca cac dong "Khong du du lieu"),
        # tu so = so dong danh dau True — dung nhu yeu cau ("tren tong so argument hop le").
        pct_significant = round(100 * sig_flags.count(True) / len(sig_flags), 1) if sig_flags else None

        pearson_r = None
        p_value = None
        if field == "overall_score" and n >= 2:
            en_scores = [r["en_avg"][field] for r in valid]
            vi_scores = [r["vi_avg"][field] for r in valid]
            r_val, p_val = stats.pearsonr(en_scores, vi_scores)
            pearson_r = round(r_val, 3)
            p_value = round(p_val, 4)

        summary[field] = {
            "n": n, "mean_diff_vi_minus_en": mean_diff, "mae": mae,
            "pearson_r": pearson_r, "p_value": p_value,
            "mean_range_en": mean_range_en, "mean_range_vi": mean_range_vi,
            "pct_significant": pct_significant,
        }

    return summary


def _direction_note(mean_diff: float | None, threshold: float = 0.2) -> str:
    if mean_diff is None:
        return "Khong co du lieu"
    if abs(mean_diff) < threshold:
        return f"Khong chenh lech dang ke (mean diff = {mean_diff:+.2f})"
    if mean_diff > 0:
        return f"Tieng Viet trung binh CAO HON tieng Anh {mean_diff:.2f} diem"
    return f"Tieng Viet trung binh THAP HON tieng Anh {abs(mean_diff):.2f} diem"


FONT = "Arial"
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def _write_header(ws, headers: list[str]):
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.freeze_panes = "A2"


def export_to_excel(rows: list[dict], summary: dict, path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Language Consistency"

    field_headers = []
    for field in AI_FIELDS:
        field_headers += [
            f"{field} EN avg", f"{field} EN range",
            f"{field} VI avg", f"{field} VI range",
            f"{field} Diff (VI-EN)", f"{field} Lech co dang ke",
        ]

    headers = (
        ["Arg ID", "Motion (EN)", "Motion (VI)", "Argument (EN)", "Argument (VI)"]
        + field_headers
        + ["Overall Reasoning EN", "Overall Reasoning VI", "Error"]
    )
    _write_header(ws, headers)

    for row_idx, r in enumerate(rows, start=2):
        col = 1
        ws.cell(row=row_idx, column=col, value=r.get("id")); col += 1
        ws.cell(row=row_idx, column=col, value=r.get("motion_en")); col += 1
        ws.cell(row=row_idx, column=col, value=r.get("motion_vi")); col += 1
        # KHONG rut gon argument text — nguoi dung can doc TOAN BO ca 2 ban de
        # tu danh gia chat luong dich truoc khi tin vao ket qua so sanh diem.
        ws.cell(row=row_idx, column=col, value=r.get("text_en")); col += 1
        ws.cell(row=row_idx, column=col, value=r.get("argument_vi")); col += 1

        en_avg = r.get("en_avg") or {}
        en_range = r.get("en_range") or {}
        vi_avg = r.get("vi_avg") or {}
        vi_range = r.get("vi_range") or {}
        diffs = r.get("diff") or {}
        significant = r.get("significant") or {}
        for field in AI_FIELDS:
            ws.cell(row=row_idx, column=col, value=en_avg.get(field)); col += 1
            ws.cell(row=row_idx, column=col, value=en_range.get(field)); col += 1
            ws.cell(row=row_idx, column=col, value=vi_avg.get(field)); col += 1
            ws.cell(row=row_idx, column=col, value=vi_range.get(field)); col += 1
            ws.cell(row=row_idx, column=col, value=diffs.get(field)); col += 1
            ws.cell(row=row_idx, column=col, value=str(significant.get(field, ""))); col += 1

        ws.cell(row=row_idx, column=col, value=r.get("en_overall_reasoning")); col += 1
        ws.cell(row=row_idx, column=col, value=r.get("vi_overall_reasoning")); col += 1
        ws.cell(row=row_idx, column=col, value=r.get("error"))

        for c in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=c).border = BORDER
            ws.cell(row=row_idx, column=c).alignment = WRAP
        ws.row_dimensions[row_idx].height = 90

    widths = [10, 30, 30, 50, 50] + [10, 10, 10, 10, 12, 14] * len(AI_FIELDS) + [40, 40, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = "Language Consistency Check — AI cham Tieng Anh vs Tieng Viet (cung noi dung)"
    ws2["A1"].font = Font(name=FONT, size=13, bold=True)

    summary_headers = [
        "Field", "N", "Mean diff (VI - EN)", "MAE",
        "Range EN avg", "Range VI avg", "% Lech dang ke",
        "Pearson r (chi Overall)", "p-value (chi Overall)",
    ]
    for col, h in enumerate(summary_headers, start=1):
        cell = ws2.cell(row=3, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER

    for i, field in enumerate(AI_FIELDS, start=4):
        m = summary.get(field, {})
        ws2.cell(row=i, column=1, value=field)
        ws2.cell(row=i, column=2, value=m.get("n"))
        ws2.cell(row=i, column=3, value=m.get("mean_diff_vi_minus_en"))
        ws2.cell(row=i, column=4, value=m.get("mae"))
        ws2.cell(row=i, column=5, value=m.get("mean_range_en"))
        ws2.cell(row=i, column=6, value=m.get("mean_range_vi"))
        ws2.cell(row=i, column=7, value=m.get("pct_significant"))
        ws2.cell(row=i, column=8, value=m.get("pearson_r"))
        ws2.cell(row=i, column=9, value=m.get("p_value"))
        for c in range(1, 10):
            ws2.cell(row=i, column=c).border = BORDER

    # Ket luan chinh, danh rieng cho overall_score (theo dung yeu cau) — dat
    # ngay duoi bang cho de thay, truoc phan chi tiet tung field.
    overall = summary.get("overall_score", {})
    conclusion_row = 4 + len(AI_FIELDS) + 1
    ws2.cell(row=conclusion_row, column=1, value="KET LUAN CHINH (overall_score):").font = Font(name=FONT, bold=True)
    ws2.cell(row=conclusion_row + 1, column=1, value="Range EN trung binh")
    ws2.cell(row=conclusion_row + 1, column=2, value=overall.get("mean_range_en"))
    ws2.cell(row=conclusion_row + 2, column=1, value="Range VI trung binh")
    ws2.cell(row=conclusion_row + 2, column=2, value=overall.get("mean_range_vi"))
    ws2.cell(row=conclusion_row + 3, column=1, value="% argument co lech dang ke (do ngon ngu, khong phai nhieu)")
    ws2.cell(row=conclusion_row + 3, column=2, value=overall.get("pct_significant"))

    note_row = conclusion_row + 5
    ws2.cell(row=note_row, column=1, value="Chieu lech (dieu huong) trung binh, tung field:").font = Font(name=FONT, bold=True)
    for i, field in enumerate(AI_FIELDS, start=note_row + 1):
        note = _direction_note(summary.get(field, {}).get("mean_diff_vi_minus_en"))
        ws2.cell(row=i, column=1, value=field).font = Font(name=FONT, bold=True)
        ws2.cell(row=i, column=2, value=note)

    read_row = note_row + len(AI_FIELDS) + 2
    reading_notes = [
        ("Cach doc ket qua:", ""),
        ("% lech dang ke THAP (vi du < 20-30%)", "Phan lon lech quan sat duoc CHI LA NHIEU NGAU NHIEN "
                                                    "(temperature), khong phai do khac ngon ngu"),
        ("% lech dang ke CAO", "Co dau hieu that su AI cham khac nhau giua 2 ngon ngu, can dieu tra prompt"),
        ("Range EN/VI trung binh cao (vi du > 1.0)", "Ban than AI da dao dong nhieu ngay ca chi doi 1 ngon ngu, "
                                                        "can canh giac khi dien giai bat ky so sanh nao khac"),
        ("r cao (> 0.7) o Overall", "2 ngon ngu xep hang cac bai theo thu tu tuong tu nhau"),
        ("QUAN TRONG", "Ket qua chi dang tin neu ban dich (cot Argument VI) dat chat luong tot — "
                        "tu kiem tra bang mat truoc khi ket luan"),
    ]
    for i, (k, v) in enumerate(reading_notes, start=read_row):
        ws2.cell(row=i, column=1, value=k).font = Font(name=FONT, bold=(v == ""))
        ws2.cell(row=i, column=2, value=v)

    ws2.column_dimensions["A"].width = 46
    ws2.column_dimensions["B"].width = 55
    for col_letter in ["C", "D", "E", "F", "G", "H", "I"]:
        ws2.column_dimensions[col_letter].width = 16

    wb.save(path)
    print(f"\nDa xuat ket qua ra: {path}")


def print_translation_samples(rows: list[dict], max_samples: int = 3) -> None:
    print("\n=== VI DU BAN DICH (de kiem tra chat luong bang mat) ===")
    shown = 0
    for r in rows:
        if not r.get("argument_vi"):
            continue
        print(f"\n--- id={r['id']} ---")
        print(f"[MOTION EN] {r['motion_en']}")
        print(f"[MOTION VI] {r['motion_vi']}")
        print(f"[ARGUMENT EN] {r['text_en'][:300]}")
        print(f"[ARGUMENT VI] {r['argument_vi'][:300]}")
        shown += 1
        if shown >= max_samples:
            break
    if shown == 0:
        print("(Khong co ban dich nao thanh cong de hien thi.)")


def main():
    # overall_reasoning + ban dich la tieng Viet co dau — tranh crash console
    # Windows giong bai hoc tu cac script truoc.
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

    parser = argparse.ArgumentParser()
    parser.add_argument("--n", type=int, default=8, help="So luong argument lay mau de test")
    parser.add_argument("--seed", type=int, default=42, help="Random seed de lay mau on dinh, lap lai duoc")
    parser.add_argument("--runs", type=int, default=2, help="So lan cham MOI ngon ngu (de do nhieu noi tai)")
    parser.add_argument(
        "--focus-id", type=str, default=None,
        help="Chi kiem tra DUNG 1 argument co #id nay (bo qua --n va --seed) — "
             "de kiem chung sau 1 case cu the, vi du '--focus-id 46238'",
    )
    args = parser.parse_args()

    settings = get_settings()
    client = get_llm_client(settings)

    if args.focus_id:
        try:
            df = load_dagstuhl_single(args.focus_id)
        except (FileNotFoundError, ValueError) as e:
            print(f"\nLOI: {e}")
            sys.exit(1)
        safe_id = "".join(c if c.isalnum() or c in "-_" else "_" for c in args.focus_id)
        output_path = Path(__file__).parent.parent / f"language_consistency_focus_{safe_id}.xlsx"
        print(f"\nCHE DO FOCUS: chi kiem tra 1 argument co id khop '{args.focus_id}' (bo qua --n va --seed).")
        print(f"Da tim thay: id={df.iloc[0]['id']}")
    else:
        df = load_dagstuhl_sample(n=args.n, seed=args.seed)
        output_path = OUTPUT_PATH

    expected_calls = len(df) * (1 + 2 * args.runs)
    print(
        f"\nDang kiem tra tinh nhat quan Anh-Viet cho {len(df)} argument x {args.runs} lan/ngon ngu "
        f"voi provider = {settings.llm_provider} ..."
    )
    print(
        f"Du kien goi API khoang {expected_calls} lan "
        f"(1 dich + {args.runs} lan cham EN + {args.runs} lan cham VI, moi argument) ...\n"
    )

    rows: list[dict] = []
    try:
        run_batch(df, settings, client, n_runs=args.runs, rows=rows)
    except KeyboardInterrupt:
        print("\n\nBi ngat boi nguoi dung (Ctrl+C) — se xuat ket qua da xu ly duoc tinh den day.")
    except Exception as e:
        print(f"\n\nDung giua chung do loi khong luong truoc: {e}")
        print("Van se xuat ket qua da xu ly duoc tinh den day.")
    finally:
        if rows:
            summary = compute_summary_metrics(rows)
            print("\n=== KET QUA LANGUAGE CONSISTENCY (EN vs VI) ===")
            print(json.dumps(summary, indent=2, ensure_ascii=False))
            export_to_excel(rows, summary, output_path)
            print_translation_samples(rows, max_samples=3)
        else:
            print("\nChua xu ly duoc argument nao, khong co gi de xuat.")


if __name__ == "__main__":
    main()
