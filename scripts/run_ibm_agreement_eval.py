"""
Chay bo dataset IBM Project Debater (opening speeches, human-authored, cham diem
that boi 15+ annotator) qua AI Evaluator, roi tinh do tuong quan (Pearson r) va
MAE giua diem AI va diem nguoi that — day la buoc "agreement rate" chinh thuc.

Moi bai duoc cham NHIEU LAN (mac dinh 2 lan, --runs de doi) de do do dao dong
giua cac lan cham cung 1 bai (do temperature > 0, khong phai do "AI doi y kien").

Nguon du lieu: https://huggingface.co/datasets/ibm-research/debate_speeches
(CDLA-Permissive-2.0 license, cong khai, khong can dang nhap)

Cach dung:
    python scripts/run_ibm_agreement_eval.py --n 20 --runs 2

Yeu cau: da dien LLM key that trong .env. Can cai them: pandas, pyarrow, requests
    pip install pandas pyarrow requests scipy

QUAN TRONG: neu bi loi/ngat giua chung (rate limit het luot retry, Ctrl+C...),
script VAN XUAT ra file Excel voi nhung bai da cham duoc, khong mat trang.
"""
import argparse
import ast
import json
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.config import get_settings
from app.schemas import ArgumentEvaluationRequest, DebateSide, DebateStage
from app.services.evaluator import evaluate_argument
from app.services.llm_client import get_llm_client

PARQUET_URL = (
    "https://huggingface.co/datasets/ibm-research/debate_speeches/"
    "resolve/refs%2Fconvert%2Fparquet/opening_speeches/train/0000.parquet"
)
LOCAL_PARQUET_PATH = Path(__file__).parent.parent / "data" / "ibm_opening_speeches.parquet"
OUTPUT_PATH = Path(__file__).parent.parent / "ibm_agreement_results.xlsx"
DELAY_BETWEEN_CALLS_SEC = 1.5
CRITERIA_NAMES = ["Logic", "Evidence", "Relevance", "Structure", "Persuasiveness"]


def download_dataset_if_needed():
    if LOCAL_PARQUET_PATH.exists():
        print(f"Da co san file: {LOCAL_PARQUET_PATH}, bo qua tai lai.")
        return
    LOCAL_PARQUET_PATH.parent.mkdir(parents=True, exist_ok=True)
    print(f"Dang tai dataset tu {PARQUET_URL} ...")
    r = requests.get(PARQUET_URL, timeout=60)
    r.raise_for_status()
    LOCAL_PARQUET_PATH.write_bytes(r.content)
    print(f"Da tai xong: {LOCAL_PARQUET_PATH} ({len(r.content) / 1024:.0f} KB)")


def load_human_expert_speeches(n: int, seed: int = 42) -> pd.DataFrame:
    """
    Chi lay source == 'Human expert' (bai phat bieu THAT cua debater chuyen gia).
    KHONG lay 'Mixed stance control (Human expert)' — do la cau hoi kiem tra
    chat luong annotator (tron 2 bai Pro+Con lam 1), khong phai bai that.
    """
    df = pd.read_parquet(LOCAL_PARQUET_PATH)
    human = df[df["source"] == "Human expert"].copy()
    if human.empty:
        raise RuntimeError(
            "Khong tim thay dong nao voi source == 'Human expert'. "
            "Kiem tra lai ten cot 'source' trong file that (co the khac ten)."
        )
    # 'goodopeningspeech' duoc luu dang chuoi kieu "[4, 5, 5, ...]" trong file
    # parquet goc, khong phai list that — can parse lai truoc khi tinh trung binh.
    human["goodopeningspeech"] = human["goodopeningspeech"].apply(
        lambda v: ast.literal_eval(v) if isinstance(v, str) else list(v)
    )
    n = min(n, len(human))
    return human.sample(n=n, random_state=seed).reset_index(drop=True)


def human_score_to_10(likert_scores: list[int]) -> float:
    """
    Likert 1-5 (trung binh nhieu annotator) -> quy doi tuyen tinh ve thang 0-10.
    Quy doi tuyen tinh KHONG lam thay doi Pearson correlation (r bat bien voi
    phep bien doi tuyen tinh) — chi phuc vu hien thi/so sanh truc quan.
    """
    avg_likert = sum(likert_scores) / len(likert_scores)
    return round((avg_likert - 1) / 4 * 10, 2)


def _evaluate_one_row(row, settings, client, n_runs: int) -> dict:
    """Cham 1 bai phat bieu N_RUNS lan doc lap. Loi o 1 lan chay KHONG lam
    hong cac lan chay khac cua CUNG bai nay (moi lan co try/except rieng)."""
    human_avg_likert = sum(row["goodopeningspeech"]) / len(row["goodopeningspeech"])
    human_score_10 = human_score_to_10(row["goodopeningspeech"])

    entry = {
        "topic_id": row["topic_id"],
        "topic": row["topic"],
        "text": row["text"],
        "n_annotators": len(row["goodopeningspeech"]),
        "human_avg_likert_1_5": round(human_avg_likert, 2),
        "human_score_10": human_score_10,
    }

    run_scores = []
    run_criteria_list = []
    run_reasoning_list = []
    req = ArgumentEvaluationRequest(
        motion=row["topic"],
        side=DebateSide.PRO,
        stage=DebateStage.OPENING,
        argument_text=row["text"],
        opponent_argument_text=None,
    )

    for run_idx in range(1, n_runs + 1):
        try:
            start = time.time()
            result = evaluate_argument(req, settings=settings, llm_client=client)
            elapsed = round(time.time() - start, 2)
            entry[f"ai_score_run_{run_idx}"] = result.overall_score
            entry[f"elapsed_run_{run_idx}"] = elapsed
            entry[f"error_run_{run_idx}"] = None
            run_scores.append(result.overall_score)
            run_criteria_list.append({c.name: c.score for c in result.criteria})
            run_reasoning_list.append(result.overall_reasoning)
            print(f"run{run_idx}={result.overall_score}", end=" ")
        except Exception as e:
            entry[f"ai_score_run_{run_idx}"] = None
            entry[f"elapsed_run_{run_idx}"] = None
            entry[f"error_run_{run_idx}"] = str(e)
            print(f"run{run_idx}=LOI({type(e).__name__})", end=" ")
        if run_idx < n_runs:
            time.sleep(DELAY_BETWEEN_CALLS_SEC)

    if run_scores:
        entry["ai_score_avg"] = round(sum(run_scores) / len(run_scores), 2)
        entry["ai_score_range"] = round(max(run_scores) - min(run_scores), 2)
        entry["diff"] = round(entry["ai_score_avg"] - human_score_10, 2)
        entry["ai_criteria"] = run_criteria_list[0] if run_criteria_list else {}
        entry["overall_reasoning"] = run_reasoning_list[0] if run_reasoning_list else None
    else:
        entry["ai_score_avg"] = None
        entry["ai_score_range"] = None
        entry["diff"] = None
        entry["ai_criteria"] = {}
        entry["overall_reasoning"] = None

    return entry


def _truncate(text: str | None, limit: int = 100) -> str:
    """Rut gon text de in ra console — them '...' o cuoi neu bi cat."""
    if not text:
        return ""
    if len(text) <= limit:
        return text
    return text[:limit].rstrip() + "..."


def run_batch(df: pd.DataFrame, settings, client, n_runs: int, rows: list) -> None:
    """
    QUAN TRONG: ghi ket qua truc tiep vao list `rows` duoc truyen vao (thay vi
    tao list moi roi return) — de neu co loi/Ctrl+C lam dung chuong trinh GIUA
    CHUNG, list `rows` ben ngoai (trong main()) van giu duoc nhung gi da cham
    xong tinh den thoi diem do, khong mat trang.
    """
    for i, row in df.iterrows():
        print(f"[{i + 1}/{len(df)}] topic_id={row['topic_id']} ...", end=" ")
        try:
            entry = _evaluate_one_row(row, settings, client, n_runs)
            rows.append(entry)
            reasoning_preview = _truncate(entry.get("overall_reasoning"), limit=100)
            print(
                f"| human={entry['human_score_10']} | ai_avg={entry['ai_score_avg']} "
                f"| reasoning: {reasoning_preview}"
            )
        except Exception as e:
            print(f"LOI TOAN BO DONG: {e}")
            rows.append({
                "topic_id": row.get("topic_id"), "topic": row.get("topic"),
                "text": row.get("text"), "n_annotators": None,
                "human_avg_likert_1_5": None, "human_score_10": None,
                "ai_score_avg": None, "ai_score_range": None, "diff": None,
                "ai_criteria": {}, "overall_reasoning": None, "error_row": str(e),
            })


def compute_agreement_metrics(rows: list[dict]) -> dict:
    """
    Pearson correlation (r) va MAE giua diem AI (trung binh cac lan cham) va
    diem nguoi that — dung phuong phap giong Themis (Hu et al., 2025).
    """
    from scipy import stats

    valid = [r for r in rows if r.get("ai_score_avg") is not None and r.get("human_score_10") is not None]
    if len(valid) < 2:
        return {"n": len(valid), "pearson_r": None, "p_value": None, "mae": None, "avg_run_range": None}

    human_scores = [r["human_score_10"] for r in valid]
    ai_scores = [r["ai_score_avg"] for r in valid]

    pearson_r, p_value = stats.pearsonr(human_scores, ai_scores)
    mae = sum(abs(h - a) for h, a in zip(human_scores, ai_scores)) / len(valid)

    ranges = [r["ai_score_range"] for r in valid if r.get("ai_score_range") is not None]
    avg_run_range = round(sum(ranges) / len(ranges), 2) if ranges else None

    return {
        "n": len(valid),
        "pearson_r": round(pearson_r, 3),
        "p_value": round(p_value, 4),
        "mae": round(mae, 2),
        "avg_run_range": avg_run_range,
    }


FONT = "Arial"
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def export_to_excel(rows: list[dict], metrics: dict, path: Path, n_runs: int):
    wb = Workbook()
    ws = wb.active
    ws.title = "Agreement Results"

    run_headers = []
    for i in range(1, n_runs + 1):
        run_headers.append(f"AI score run {i}")

    headers = (
        ["Topic ID", "Topic", "Text (truncated)", "# Annotators",
         "Human avg (1-5 Likert)", "Human score (0-10)"]
        + run_headers
        + ["AI score AVG", "Range (max-min giua cac lan)", "Diff (AI avg - Human)"]
        + ["Overall Reasoning (run 1)"]
        + CRITERIA_NAMES
        + ["Error"]
    )
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.freeze_panes = "A2"

    for row_idx, r in enumerate(rows, start=2):
        col = 1
        ws.cell(row=row_idx, column=col, value=r.get("topic_id")); col += 1
        ws.cell(row=row_idx, column=col, value=r.get("topic")); col += 1
        text = r.get("text") or ""
        ws.cell(row=row_idx, column=col, value=text[:300]); col += 1
        ws.cell(row=row_idx, column=col, value=r.get("n_annotators")); col += 1
        ws.cell(row=row_idx, column=col, value=r.get("human_avg_likert_1_5")); col += 1
        ws.cell(row=row_idx, column=col, value=r.get("human_score_10")); col += 1
        for i in range(1, n_runs + 1):
            ws.cell(row=row_idx, column=col, value=r.get(f"ai_score_run_{i}")); col += 1
        ws.cell(row=row_idx, column=col, value=r.get("ai_score_avg")); col += 1
        ws.cell(row=row_idx, column=col, value=r.get("ai_score_range")); col += 1
        ws.cell(row=row_idx, column=col, value=r.get("diff")); col += 1
        ws.cell(row=row_idx, column=col, value=r.get("overall_reasoning")); col += 1
        for cname in CRITERIA_NAMES:
            ws.cell(row=row_idx, column=col, value=(r.get("ai_criteria") or {}).get(cname)); col += 1
        error_text = r.get("error_row") or "; ".join(
            f"run{i}: {r.get(f'error_run_{i}')}" for i in range(1, n_runs + 1) if r.get(f"error_run_{i}")
        )
        ws.cell(row=row_idx, column=col, value=error_text or None)

        for c in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=c).border = BORDER
            ws.cell(row=row_idx, column=c).alignment = WRAP
        ws.row_dimensions[row_idx].height = 50

    widths = [10, 30, 44, 12, 16, 14] + [12] * n_runs + [12, 16, 16] + [50] + [10] * 5 + [30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("Agreement Metrics")
    ws2["A1"] = "Agreement Rate — AI vs Human (IBM Project Debater dataset)"
    ws2["A1"].font = Font(name=FONT, size=13, bold=True)
    info = [
        ("N (so cap diem hop le)", metrics["n"]),
        ("Pearson correlation (r)", metrics["pearson_r"]),
        ("p-value", metrics["p_value"]),
        ("MAE (Mean Absolute Error, thang 0-10)", metrics["mae"]),
        ("Do dao dong TB giua cac lan cham (range)", metrics["avg_run_range"]),
        ("", ""),
        ("Cach doc ket qua:", ""),
        ("r > 0.5", "Tuong quan kha, AI va nguoi co xu huong danh gia giong nhau"),
        ("r 0.3 - 0.5", "Tuong quan yeu, can xem lai rubric/prompt"),
        ("r < 0.3", "Gan nhu khong tuong quan, AI dang danh gia khac han nguoi"),
        ("p-value < 0.05", "Ket qua co y nghia thong ke (khong phai ngau nhien)"),
        ("Range trung binh cao (vi du > 1.0)", "Diem dao dong nhieu giua cac lan cham CUNG 1 bai -> "
                                                "1 phan nguyen nhan MAE/r xau la do nhieu ngau nhien "
                                                "(temperature), khong chi do rubric qua khat khe"),
    ]
    for i, (k, v) in enumerate(info, start=3):
        ws2.cell(row=i, column=1, value=k).font = Font(name=FONT, bold=bool(k) and ":" not in k)
        ws2.cell(row=i, column=2, value=v)
    ws2.column_dimensions["A"].width = 40
    ws2.column_dimensions["B"].width = 60

    wb.save(path)
    print(f"\nDa xuat ket qua ra: {path}")


def main():
    # overall_reasoning tu LLM la tieng Viet co dau — console Windows mac dinh dung
    # codepage (vi du cp1252) khong encode duoc, se crash print() giua chung va lam
    # loi cai chinh xac VUA cham xong bi ghi de thanh 1 dong loi trung lap (xem
    # run_batch: print that bai -> roi vao except -> append them 1 dong "LOI" du
    # entry dung da duoc them truoc do). Ep stdout sang UTF-8, thay ky tu khong
    # encode duoc bang '?' thay vi crash, de tranh corrupt du lieu ket qua.
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

    parser = argparse.ArgumentParser()
    parser.add_argument("--n", type=int, default=20, help="So luong bai phat bieu lay mau de test")
    parser.add_argument("--seed", type=int, default=42, help="Random seed de lay mau on dinh, lap lai duoc")
    parser.add_argument("--runs", type=int, default=2, help="So lan cham MOI bai (de do dao dong)")
    args = parser.parse_args()

    download_dataset_if_needed()

    settings = get_settings()
    client = get_llm_client(settings)

    df = load_human_expert_speeches(n=args.n, seed=args.seed)
    print(
        f"\nDang cham {len(df)} bai phat bieu 'Human expert' x {args.runs} lan/bai "
        f"voi provider = {settings.llm_provider} ...\n"
    )

    rows: list[dict] = []
    try:
        run_batch(df, settings, client, n_runs=args.runs, rows=rows)
    except KeyboardInterrupt:
        print("\n\nBi ngat boi nguoi dung (Ctrl+C) — se xuat ket qua da cham duoc tinh den day.")
    except Exception as e:
        print(f"\n\nDung giua chung do loi khong luong truoc: {e}")
        print("Van se xuat ket qua da cham duoc tinh den day.")
    finally:
        if rows:
            metrics = compute_agreement_metrics(rows)
            print("\n=== KET QUA AGREEMENT (tren so bai da cham xong) ===")
            print(json.dumps(metrics, indent=2, ensure_ascii=False))
            export_to_excel(rows, metrics, OUTPUT_PATH, n_runs=args.runs)
        else:
            print("\nChua cham duoc bai nao, khong co gi de xuat.")


if __name__ == "__main__":
    main()
