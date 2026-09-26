"""
Kiem tra DO DUNG DAN cua reasoning (KHONG phai do lech diem so) khi cham CUNG
1 bai NHIEU LAN bang tieng Viet.

BOI CANH: case 46238 (Dagstuhl-15512) cham 5 lan cho range diem 2-3, nhung doc
qua thi overall_reasoning giua cac lan TRONG kha giong nhau ve mat nhan xet
dinh tinh (deu noi "logic yeu, thieu bang chung, cau truc lon xon" hay tuong
duong). Nghi ngo: AI co the HIEU DUNG va NHAT QUAN o tang nhan xet dinh tinh,
chi "lam tron" thanh so khac nhau moi lan — day la van de NHO HON nhieu so
voi viec AI hieu sai/mau thuan giua cac lan. Script nay kiem tra dieu do RONG
hon (nhieu argument, khong chi 1 case) va DAY DU hon (luu toan bo text
reasoning cua MOI lan chay, khong chi 1 lan dai dien).

Khac voi run_language_consistency_check.py: script nay CHI cham tieng Viet
(khong can ban tieng Anh doi chieu), KHONG tinh diff/Pearson r/MAE gi ca — day
la cong cu DINH TINH, doc bang mat, khong phai dinh luong.

Cach dung:
    python scripts/run_reasoning_consistency_check.py --n 5 --runs 3 --seed 42

Yeu cau: da dien LLM key that trong .env, va da co san
data/dagstuhl_processed.csv (chay scripts/run_dagstuhl_agreement_eval.py it
nhat 1 lan truoc de tao file cache nay).

QUAN TRONG: giong cac script truoc, neu bi loi/ngat giua chung, script VAN
XUAT ra file Excel voi nhung argument da xu ly xong, khong mat trang.
"""
import argparse
import json
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.config import get_settings, Settings
from app.schemas import ArgumentEvaluationRequest, DebateSide, DebateStage
from app.services.evaluator import evaluate_argument
from app.services.json_utils import EvaluationParseError, call_llm_with_retry
from app.services.llm_client import LLMClient, get_llm_client

PROCESSED_CSV_PATH = Path(__file__).parent.parent / "data" / "dagstuhl_processed.csv"
OUTPUT_PATH = Path(__file__).parent.parent / "reasoning_consistency_results.xlsx"
DELAY_BETWEEN_CALLS_SEC = 1.5

CRITERIA_NAMES = ["Logic", "Evidence", "Relevance", "Structure", "Persuasiveness"]


# --- Dich sang tieng Viet: tai su dung dung logic (khong dung lai code) tu
# run_language_consistency_check.py — prompt dich don gian, output JSON 2 field. ---

TRANSLATE_SYSTEM_PROMPT = """Ban la 1 bien dich vien chuyen nghiep, dich noi dung tranh bien tu tieng Anh \
sang tieng Viet.

NHIEM VU: Dich CHINH XAC va tu nhien, GIU NGUYEN y nghia, sac thai, va muc do \
manh/yeu cua lap luan trong ban goc — KHONG dien giai them, KHONG tom tat, \
KHONG bo sot y, KHONG them binh luan hay sua loi lap luan cua ban goc (ke ca \
neu ban goc co lap luan yeu hoac sai — dich dung y, khong "sua ho").

Dich CA HAI: "motion" (chu de tranh bien) va "argument" (noi dung lap luan).

[OUTPUT - CHI tra ve JSON hop le, khong them text nao khac ngoai JSON]
{
  "motion_vi": "<ban dich tieng Viet cua motion>",
  "argument_vi": "<ban dich tieng Viet cua argument, giu nguyen so doan/cau truc neu co>"
}"""


def build_translate_user_prompt(motion: str, argument_text: str) -> str:
    return f"[MOTION]\n{motion}\n\n[ARGUMENT]\n{argument_text}"


def _build_translation_result(parsed: dict, raw_text: str) -> tuple[str, str]:
    motion_vi = parsed.get("motion_vi")
    argument_vi = parsed.get("argument_vi")
    if not motion_vi or not argument_vi:
        raise EvaluationParseError(
            "LLM khong tra ve du 'motion_vi' va 'argument_vi' khi dich — coi la loi dinh dang de retry."
        )
    return motion_vi, argument_vi


def translate_to_vietnamese(motion: str, argument_text: str, client: LLMClient) -> tuple[str, str]:
    user_prompt = build_translate_user_prompt(motion, argument_text)
    return call_llm_with_retry(
        generate_fn=lambda up: client.generate(TRANSLATE_SYSTEM_PROMPT, up),
        system_prompt=TRANSLATE_SYSTEM_PROMPT,
        user_prompt=user_prompt,
        build_result_fn=_build_translation_result,
    )


def load_dagstuhl_sample(n: int, seed: int = 42) -> pd.DataFrame:
    if not PROCESSED_CSV_PATH.exists():
        raise FileNotFoundError(
            f"Khong tim thay {PROCESSED_CSV_PATH}. Chay scripts/run_dagstuhl_agreement_eval.py "
            "it nhat 1 lan truoc de tao file cache nay."
        )
    df = pd.read_csv(PROCESSED_CSV_PATH)
    n = min(n, len(df))
    return df.sample(n=n, random_state=seed).reset_index(drop=True)


def _evaluate_one_row(row, settings: Settings, client: LLMClient, n_runs: int) -> dict:
    """
    Dich 1 lan, roi cham ban tieng Viet do `n_runs` lan DOC LAP. Luu DAY DU
    (khong rut gon) overall_score/overall_reasoning + 5 tieu chi (score +
    reasoning) cua TUNG lan chay rieng biet trong entry["runs"] (list).
    """
    entry = {
        "id": row["id"],
        "motion_en": row["issue"],
        "text_en": row["text"],
        "motion_vi": None,
        "argument_vi": None,
        "runs": [],
        "error": None,
    }

    try:
        motion_vi, argument_vi = translate_to_vietnamese(entry["motion_en"], entry["text_en"], client)
        entry["motion_vi"] = motion_vi
        entry["argument_vi"] = argument_vi
    except Exception as e:
        entry["error"] = f"Dich loi: {e}"
        return entry
    time.sleep(DELAY_BETWEEN_CALLS_SEC)

    req = ArgumentEvaluationRequest(
        motion=entry["motion_vi"],
        side=DebateSide.PRO,  # context SOLO -> side khong anh huong ban chat cham diem
        stage=DebateStage.OPENING,
        argument_text=entry["argument_vi"],
        opponent_argument_text=None,
    )

    n_success = 0
    for run_idx in range(1, n_runs + 1):
        try:
            result = evaluate_argument(req, settings=settings, llm_client=client)
            run_entry = {
                "run": run_idx,
                "overall_score": result.overall_score,
                "overall_reasoning": result.overall_reasoning,
                "criteria": {c.name: {"score": c.score, "reasoning": c.reasoning} for c in result.criteria},
                "error": None,
            }
            n_success += 1
            print(f"run{run_idx}=OK(score={result.overall_score})", end=" ")
        except Exception as e:
            run_entry = {
                "run": run_idx, "overall_score": None, "overall_reasoning": None,
                "criteria": {}, "error": str(e),
            }
            print(f"run{run_idx}=LOI({type(e).__name__})", end=" ")
        entry["runs"].append(run_entry)
        if run_idx < n_runs:
            time.sleep(DELAY_BETWEEN_CALLS_SEC)

    if n_success == 0:
        entry["error"] = "Cham tieng Viet that bai toan bo cac lan chay"

    return entry


def run_batch(df: pd.DataFrame, settings: Settings, client: LLMClient, n_runs: int, rows: list) -> None:
    """Ghi truc tiep vao `rows` truyen vao — giong cac script truoc, de neu
    loi/Ctrl+C giua chung, ket qua da xu ly duoc khong bi mat."""
    for i, row in df.iterrows():
        print(f"[{i + 1}/{len(df)}] id={row['id']} ...", end=" ")
        try:
            entry = _evaluate_one_row(row, settings, client, n_runs)
            rows.append(entry)
            print("| DONE" if not entry["error"] else f"\n    LOI: {entry['error']}")
        except Exception as e:
            print(f"LOI TOAN BO DONG: {e}")
            rows.append({
                "id": row.get("id"), "motion_en": row.get("issue"), "text_en": row.get("text"),
                "motion_vi": None, "argument_vi": None, "runs": [], "error": str(e),
            })


FONT = "Arial"
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
# 2 mau xen ke theo tung NHOM argument (khong phai tung dong) — de mat de theo
# doi cac dong lien tiep cua CUNG 1 bai khi doc trong Excel.
GROUP_FILL_A = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
GROUP_FILL_B = PatternFill(start_color="F2F6FA", end_color="F2F6FA", fill_type="solid")


def _write_header(ws, headers: list[str]):
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.freeze_panes = "A2"


def export_to_excel(rows: list[dict], path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Reasoning Detail"

    criteria_headers = []
    for cname in CRITERIA_NAMES:
        criteria_headers += [f"{cname} Score", f"{cname} Reasoning"]

    headers = (
        ["Arg ID", "Run", "Motion (VI)", "Argument (VI)", "Overall Score", "Overall Reasoning"]
        + criteria_headers
        + ["Run Error", "Nhan xet dung khong (Y/N)", "Ghi chu"]
    )
    _write_header(ws, headers)

    row_idx = 2
    group_fill_toggle = False
    for r in rows:
        group_fill = GROUP_FILL_B if group_fill_toggle else GROUP_FILL_A
        group_fill_toggle = not group_fill_toggle

        run_list = r.get("runs") or []
        if not run_list:
            # Ca argument nay loi ngay tu buoc dich (khong co lan chay nao) —
            # van ghi 1 dong de khong mat dau vet, kem loi ro rang.
            ws.cell(row=row_idx, column=1, value=r.get("id"))
            ws.cell(row=row_idx, column=2, value=None)
            ws.cell(row=row_idx, column=3, value=r.get("motion_vi"))
            ws.cell(row=row_idx, column=4, value=r.get("argument_vi"))
            error_col = 6 + len(criteria_headers) + 1
            ws.cell(row=row_idx, column=error_col, value=r.get("error"))
            for c in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=c).border = BORDER
                ws.cell(row=row_idx, column=c).alignment = WRAP
                ws.cell(row=row_idx, column=c).fill = group_fill
            ws.row_dimensions[row_idx].height = 30
            row_idx += 1
            continue

        first_row_of_group = row_idx
        for run_entry in run_list:
            col = 1
            ws.cell(row=row_idx, column=col, value=r.get("id")); col += 1
            ws.cell(row=row_idx, column=col, value=run_entry.get("run")); col += 1
            # Motion/Argument (VI) CHI hien o dong DAU TIEN cua nhom, do bot lap.
            ws.cell(row=row_idx, column=col, value=r.get("motion_vi") if row_idx == first_row_of_group else None); col += 1
            ws.cell(row=row_idx, column=col, value=r.get("argument_vi") if row_idx == first_row_of_group else None); col += 1
            ws.cell(row=row_idx, column=col, value=run_entry.get("overall_score")); col += 1
            ws.cell(row=row_idx, column=col, value=run_entry.get("overall_reasoning")); col += 1

            criteria = run_entry.get("criteria") or {}
            for cname in CRITERIA_NAMES:
                c_data = criteria.get(cname) or {}
                ws.cell(row=row_idx, column=col, value=c_data.get("score")); col += 1
                ws.cell(row=row_idx, column=col, value=c_data.get("reasoning")); col += 1

            ws.cell(row=row_idx, column=col, value=run_entry.get("error")); col += 1
            # 2 cot trong de nguoi doc tu dien tay — CHU Y: KHONG ghi gia tri gi vao day.
            col += 2

            for c in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=c).border = BORDER
                ws.cell(row=row_idx, column=c).alignment = WRAP
                ws.cell(row=row_idx, column=c).fill = group_fill
            ws.row_dimensions[row_idx].height = 70
            row_idx += 1

    widths = (
        [10, 6, 30, 44, 12, 44]
        + [8, 34] * len(CRITERIA_NAMES)
        + [30, 20, 30]
    )
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Sheet phu: overall_score qua cac lan chay, canh nhau — CHI de tham khao
    # nhanh dao dong so, KHONG phai trong tam cua script nay.
    ws2 = wb.create_sheet("Overall Scores Overview")
    max_runs = max((len(r.get("runs") or []) for r in rows), default=0)
    overview_headers = ["Arg ID", "Motion (VI)"] + [f"Run {i} score" for i in range(1, max_runs + 1)] + ["Error"]
    _write_header(ws2, overview_headers)

    for i, r in enumerate(rows, start=2):
        ws2.cell(row=i, column=1, value=r.get("id"))
        ws2.cell(row=i, column=2, value=r.get("motion_vi"))
        run_list = r.get("runs") or []
        for j in range(max_runs):
            score = run_list[j]["overall_score"] if j < len(run_list) else None
            ws2.cell(row=i, column=3 + j, value=score)
        ws2.cell(row=i, column=3 + max_runs, value=r.get("error"))
        for c in range(1, len(overview_headers) + 1):
            ws2.cell(row=i, column=c).border = BORDER

    ws2.column_dimensions["A"].width = 10
    ws2.column_dimensions["B"].width = 34
    for j in range(max_runs):
        ws2.column_dimensions[get_column_letter(3 + j)].width = 12
    ws2.column_dimensions[get_column_letter(3 + max_runs)].width = 30

    wb.save(path)
    print(f"\nDa xuat ket qua ra: {path}")


def print_full_example(rows: list[dict]) -> None:
    """In DAY DU (khong rut gon) tat ca cac lan chay cua 1 argument dau tien
    xu ly THANH CONG, de xem truoc chat luong/tinh nhat quan cua reasoning
    truoc khi mo Excel."""
    for r in rows:
        run_list = r.get("runs") or []
        if not run_list or all(run["overall_score"] is None for run in run_list):
            continue

        print("\n" + "=" * 70)
        print(f"VI DU DAY DU — id={r['id']}")
        print("=" * 70)
        print(f"[MOTION VI] {r['motion_vi']}")
        print(f"[ARGUMENT VI] {r['argument_vi']}")

        for run_entry in run_list:
            print(f"\n--- Run {run_entry['run']} ---")
            if run_entry.get("error"):
                print(f"LOI: {run_entry['error']}")
                continue
            print(f"overall_score = {run_entry['overall_score']}")
            print(f"overall_reasoning: {run_entry['overall_reasoning']}")
            for cname in CRITERIA_NAMES:
                c_data = (run_entry.get("criteria") or {}).get(cname) or {}
                print(f"  [{cname}] score={c_data.get('score')} — {c_data.get('reasoning')}")
        return

    print("\n(Khong co argument nao xu ly thanh cong de in vi du.)")


def main():
    # overall_reasoning + ban dich la tieng Viet co dau — tranh crash console
    # Windows giong bai hoc tu cac script truoc.
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

    parser = argparse.ArgumentParser()
    parser.add_argument("--n", type=int, default=5, help="So luong argument lay mau de test")
    parser.add_argument("--seed", type=int, default=42, help="Random seed de lay mau on dinh, lap lai duoc")
    parser.add_argument("--runs", type=int, default=3, help="So lan cham DOC LAP ban tieng Viet cua MOI argument")
    args = parser.parse_args()

    settings = get_settings()
    client = get_llm_client(settings)

    df = load_dagstuhl_sample(n=args.n, seed=args.seed)

    expected_calls = len(df) * (1 + args.runs)
    print(
        f"\nDang kiem tra do dung dan cua reasoning (tieng Viet, --runs={args.runs}) "
        f"cho {len(df)} argument voi provider = {settings.llm_provider} ..."
    )
    print(
        f"Du kien goi API khoang {expected_calls} lan "
        f"(1 dich + {args.runs} lan cham, moi argument) ...\n"
    )

    rows: list[dict] = []
    try:
        run_batch(df, settings, client, n_runs=args.runs, rows=rows)
    except KeyboardInterrupt:
        print("\n\nBi ngat boi nguoi dung (Ctrl+C) — se xuat ket qua da xu ly duoc tinh den day.")
    except Exception as e:
        print(f"\n\nDung giua chung do loi khong luong truoc: {e}")
        print("Van se xuat ket qua da xu ly duoc tinh den day.")
    finally:
        if rows:
            export_to_excel(rows, OUTPUT_PATH)
            print_full_example(rows)
        else:
            print("\nChua xu ly duoc argument nao, khong co gi de xuat.")


if __name__ == "__main__":
    main()
