"""
Chay bo test case Session (nhieu Round moi session) qua AI Evaluator THAT,
xuat ket qua ra Excel 3 sheet: Round Results, Session Results, Summary.

Cach dung:
    python scripts/run_session_batch_eval.py

Yeu cau: da dien LLM key that trong .env. Script nay goi API THAT theo dung
so luong round + session trong data/session_test_cases.json (moi session co
N round -> N lan goi /evaluate/round + 1 lan goi /evaluate/session).
"""
import json
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.config import get_settings, Settings
from app.schemas import (
    ArgumentTurn,
    DebateSide,
    DebateStage,
    RoundEvaluationRequest,
    RoundEvaluationResult,
    SessionEvaluationRequest,
    TurnSpeaker,
)
from app.services.round_evaluator import evaluate_round
from app.services.session_evaluator import evaluate_session
from app.services.llm_client import LLMClient, get_llm_client

DATA_PATH = Path(__file__).parent.parent / "data" / "session_test_cases.json"
OUTPUT_PATH = Path(__file__).parent.parent / "session_batch_results.xlsx"
DELAY_BETWEEN_CALLS_SEC = 1.5

CRITERIA_NAMES = ["Logic", "Evidence", "Relevance", "Structure", "Persuasiveness"]


def load_sessions(path: Path) -> list[dict]:
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def _to_round_request(session: dict, round_data: dict) -> RoundEvaluationRequest:
    turns = [
        ArgumentTurn(speaker=TurnSpeaker(t["speaker"]), text=t["text"])
        for t in round_data["turns"]
    ]
    return RoundEvaluationRequest(
        motion=session["motion"],
        side=DebateSide(session["side"]),
        stage=DebateStage(round_data["stage"]),
        turns=turns,
    )


def run_all_sessions(
    sessions: list[dict], settings: Settings, client: LLMClient
) -> tuple[list[dict], list[dict]]:
    round_rows: list[dict] = []
    session_rows: list[dict] = []

    for s_idx, session in enumerate(sessions, start=1):
        print(f"\n[{s_idx}/{len(sessions)}] Session: {session['id']} ({session['category']})")
        round_results: list[RoundEvaluationResult] = []
        session_error = None

        for r_idx, round_data in enumerate(session["rounds"], start=1):
            stage = round_data["stage"]
            print(f"    [Round {r_idx}] stage={stage} ...", end=" ")
            try:
                req = _to_round_request(session, round_data)
                start = time.time()
                result = evaluate_round(req, settings=settings, llm_client=client)
                elapsed = round(time.time() - start, 2)
                round_results.append(result)
                round_rows.append({
                    "session_id": session["id"],
                    "category": session["category"],
                    "stage": stage,
                    "round_score": result.round_score,
                    "criteria": {c.name: (c.score, c.reasoning) for c in result.criteria},
                    "consistency_note": result.consistency_note,
                    "strengths": result.strengths,
                    "weaknesses": result.weaknesses,
                    "suggestions": result.suggestions,
                    "elapsed_sec": elapsed,
                    "error": None,
                })
                print(f"round_score={result.round_score}")
            except Exception as e:
                print(f"LOI: {e}")
                round_rows.append({
                    "session_id": session["id"],
                    "category": session["category"],
                    "stage": stage,
                    "round_score": None,
                    "criteria": {},
                    "consistency_note": None,
                    "strengths": [],
                    "weaknesses": [],
                    "suggestions": [],
                    "elapsed_sec": None,
                    "error": str(e),
                })
                session_error = f"Round '{stage}' loi: {e}"
            time.sleep(DELAY_BETWEEN_CALLS_SEC)

        if not round_results:
            session_rows.append({
                "session_id": session["id"], "category": session["category"],
                "overall_score": None, "stage_breakdown": {}, "progress_trend": None,
                "overall_strengths": [], "overall_weaknesses": [], "overall_suggestions": [],
                "error": session_error or "Khong co round nao thanh cong",
            })
            continue

        print("    [Session synthesis] ...", end=" ")
        try:
            session_req = SessionEvaluationRequest(
                motion=session["motion"],
                side=DebateSide(session["side"]),
                rounds=round_results,
            )
            session_result = evaluate_session(session_req, settings=settings, llm_client=client)
            session_rows.append({
                "session_id": session["id"],
                "category": session["category"],
                "overall_score": session_result.overall_score,
                "stage_breakdown": session_result.stage_breakdown,
                "progress_trend": session_result.progress_trend,
                "overall_strengths": session_result.overall_strengths,
                "overall_weaknesses": session_result.overall_weaknesses,
                "overall_suggestions": session_result.overall_suggestions,
                "error": session_error,
            })
            print(f"overall_score={session_result.overall_score}")
        except Exception as e:
            print(f"LOI: {e}")
            session_rows.append({
                "session_id": session["id"], "category": session["category"],
                "overall_score": None, "stage_breakdown": {}, "progress_trend": None,
                "overall_strengths": [], "overall_weaknesses": [], "overall_suggestions": [],
                "error": str(e),
            })
        time.sleep(DELAY_BETWEEN_CALLS_SEC)

    return round_rows, session_rows


FONT = "Arial"
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def _write_header(ws, headers: list[str]):
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.freeze_panes = "A2"


def export_round_sheet(wb, round_rows: list[dict]):
    ws = wb.active
    ws.title = "Round Results"

    headers = (
        ["Session ID", "Category", "Stage", "Round score (0-10)"]
        + [f"{c} (1-5)" for c in CRITERIA_NAMES]
        + [f"{c} reasoning" for c in CRITERIA_NAMES]
        + ["Consistency note", "Strengths", "Weaknesses", "Suggestions",
           "Time (s)", "Diem nguoi cham round (0-10)", "Ghi chu", "Error"]
    )
    _write_header(ws, headers)

    for row_idx, r in enumerate(round_rows, start=2):
        col = 1
        ws.cell(row=row_idx, column=col, value=r["session_id"]); col += 1
        ws.cell(row=row_idx, column=col, value=r["category"]); col += 1
        ws.cell(row=row_idx, column=col, value=r["stage"]); col += 1
        ws.cell(row=row_idx, column=col, value=r["round_score"]); col += 1
        for cname in CRITERIA_NAMES:
            score, _ = r["criteria"].get(cname, (None, ""))
            ws.cell(row=row_idx, column=col, value=score); col += 1
        for cname in CRITERIA_NAMES:
            _, reasoning = r["criteria"].get(cname, (None, ""))
            ws.cell(row=row_idx, column=col, value=reasoning); col += 1
        ws.cell(row=row_idx, column=col, value=r["consistency_note"]); col += 1
        ws.cell(row=row_idx, column=col, value="; ".join(r["strengths"])); col += 1
        ws.cell(row=row_idx, column=col, value="; ".join(r["weaknesses"])); col += 1
        ws.cell(row=row_idx, column=col, value="; ".join(r["suggestions"])); col += 1
        ws.cell(row=row_idx, column=col, value=r["elapsed_sec"]); col += 1
        col += 1
        col += 1
        ws.cell(row=row_idx, column=col, value=r["error"])

        for c in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=c).border = BORDER
            ws.cell(row=row_idx, column=c).alignment = WRAP
        ws.row_dimensions[row_idx].height = 60

    widths = [20, 26, 12, 14] + [10] * 5 + [32] * 5 + [34, 30, 30, 30, 10, 20, 24, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def export_session_sheet(wb, session_rows: list[dict]):
    ws = wb.create_sheet("Session Results")
    headers = [
        "Session ID", "Category", "Overall score (0-10)",
        "Opening", "Rebuttal", "Closing",
        "Progress trend", "Overall strengths", "Overall weaknesses", "Overall suggestions",
        "Diem nguoi cham session (0-10)", "Ghi chu", "Error",
    ]
    _write_header(ws, headers)

    for row_idx, r in enumerate(session_rows, start=2):
        sb = r["stage_breakdown"] or {}
        col = 1
        ws.cell(row=row_idx, column=col, value=r["session_id"]); col += 1
        ws.cell(row=row_idx, column=col, value=r["category"]); col += 1
        ws.cell(row=row_idx, column=col, value=r["overall_score"]); col += 1
        ws.cell(row=row_idx, column=col, value=sb.get("opening")); col += 1
        ws.cell(row=row_idx, column=col, value=sb.get("rebuttal")); col += 1
        ws.cell(row=row_idx, column=col, value=sb.get("closing")); col += 1
        ws.cell(row=row_idx, column=col, value=r["progress_trend"]); col += 1
        ws.cell(row=row_idx, column=col, value="; ".join(r["overall_strengths"])); col += 1
        ws.cell(row=row_idx, column=col, value="; ".join(r["overall_weaknesses"])); col += 1
        ws.cell(row=row_idx, column=col, value="; ".join(r["overall_suggestions"])); col += 1
        col += 1
        col += 1
        ws.cell(row=row_idx, column=col, value=r["error"])

        for c in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=c).border = BORDER
            ws.cell(row=row_idx, column=c).alignment = WRAP
        ws.row_dimensions[row_idx].height = 70

    widths = [20, 26, 16, 10, 10, 10, 40, 34, 34, 34, 22, 24, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def export_summary_sheet(wb, round_rows: list[dict], session_rows: list[dict]):
    ws = wb.create_sheet("Summary")
    ws["A1"] = "Session Batch Evaluation Summary"
    ws["A1"].font = Font(name=FONT, size=13, bold=True)

    round_scores = [r["round_score"] for r in round_rows if r["round_score"] is not None]
    session_scores = [r["overall_score"] for r in session_rows if r["overall_score"] is not None]

    summary = [
        ("Total sessions", len(session_rows)),
        ("Total rounds", len(round_rows)),
        ("Rounds succeeded", len(round_scores)),
        ("Rounds failed", len(round_rows) - len(round_scores)),
        ("Sessions succeeded", len(session_scores)),
        ("Sessions failed", len(session_rows) - len(session_scores)),
        ("Average round_score", round(sum(round_scores) / len(round_scores), 2) if round_scores else "-"),
        ("Average session overall_score", round(sum(session_scores) / len(session_scores), 2) if session_scores else "-"),
    ]
    for i, (k, v) in enumerate(summary, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(name=FONT, bold=True)
        ws.cell(row=i, column=2, value=v)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 18


def export_to_excel(round_rows: list[dict], session_rows: list[dict], path: Path):
    wb = Workbook()
    export_round_sheet(wb, round_rows)
    export_session_sheet(wb, session_rows)
    export_summary_sheet(wb, round_rows, session_rows)
    wb.save(path)
    print(f"\nDa xuat ket qua ra: {path}")


def main():
    settings = get_settings()
    client = get_llm_client(settings)
    sessions = load_sessions(DATA_PATH)
    print(f"Dang chay {len(sessions)} session voi provider = {settings.llm_provider} ...")
    round_rows, session_rows = run_all_sessions(sessions, settings, client)
    export_to_excel(round_rows, session_rows, OUTPUT_PATH)


if __name__ == "__main__":
    main()
