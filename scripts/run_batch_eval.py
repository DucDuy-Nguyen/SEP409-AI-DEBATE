"""
Chay mot bo test case debate qua AI Evaluator THAT (goi API that, ton chi phi/quota),
xuat ket qua ra file Excel de doi chieu voi diem nguoi cham that (agreement rate).

Cach dung:
    python scripts/run_batch_eval.py

Yeu cau: da dien LLM key that trong .env (KHONG dung FakeLLMClient — day la
batch that, moi lan chay se goi API that theo so luong test case trong
data/test_cases.json).
"""
import json
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.config import get_settings
from app.schemas import ArgumentEvaluationRequest, DebateSide, DebateStage
from app.services.evaluator import evaluate_argument, EvaluationParseError
from app.services.llm_client import get_llm_client

DATA_PATH = Path(__file__).parent.parent / "data" / "test_cases.json"
OUTPUT_PATH = Path(__file__).parent.parent / "batch_eval_results.xlsx"
DELAY_BETWEEN_CALLS_SEC = 1.5  # tranh bi rate-limit tren free tier


def load_test_cases():
    with open(DATA_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


def run_all_cases(cases, settings, client):
    results = []
    for i, case in enumerate(cases, start=1):
        print(f"[{i}/{len(cases)}] Dang cham: {case['id']} ...")
        req = ArgumentEvaluationRequest(
            motion=case["motion"],
            side=DebateSide(case["side"]),
            stage=DebateStage(case.get("stage", "rebuttal")),
            argument_text=case["argument_text"],
            opponent_argument_text=case.get("opponent_argument_text"),
        )
        start = time.time()
        try:
            result = evaluate_argument(req, settings=settings, llm_client=client)
            elapsed = round(time.time() - start, 2)
            results.append({
                "id": case["id"],
                "category": case.get("category", ""),
                "expected_tier": case.get("expected_tier", ""),
                "motion": case["motion"],
                "argument_text": case["argument_text"],
                "overall_score": result.overall_score,
                "criteria": {c.name: (c.score, c.reasoning) for c in result.criteria},
                "strengths": result.strengths,
                "weaknesses": result.weaknesses,
                "suggestions": result.suggestions,
                "elapsed_sec": elapsed,
                "error": None,
            })
            print(f"    -> overall_score = {result.overall_score}")
        except Exception as e:
            results.append({
                "id": case["id"],
                "category": case.get("category", ""),
                "expected_tier": case.get("expected_tier", ""),
                "motion": case["motion"],
                "argument_text": case["argument_text"],
                "overall_score": None,
                "criteria": {},
                "strengths": [],
                "weaknesses": [],
                "suggestions": [],
                "elapsed_sec": None,
                "error": str(e),
            })
            print(f"    -> LOI: {e}")
        time.sleep(DELAY_BETWEEN_CALLS_SEC)
    return results


def export_to_excel(results, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Batch Results"

    FONT = "Arial"
    HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    THIN = Side(style="thin", color="BFBFBF")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    WRAP = Alignment(wrap_text=True, vertical="top")

    criteria_names = ["Logic", "Evidence", "Relevance", "Structure", "Persuasiveness"]

    headers = (
        ["Test ID", "Category", "Expected tier", "Motion", "Argument text", "Overall score (0-10)"]
        + [f"{c} (1-5)" for c in criteria_names]
        + [f"{c} reasoning" for c in criteria_names]
        + ["Strengths", "Weaknesses", "Suggestions", "Time (s)",
           "Diem nguoi cham (0-10)", "Ghi chu", "Error"]
    )

    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    for row_idx, r in enumerate(results, start=2):
        col = 1
        ws.cell(row=row_idx, column=col, value=r["id"]); col += 1
        ws.cell(row=row_idx, column=col, value=r["category"]); col += 1
        ws.cell(row=row_idx, column=col, value=r["expected_tier"]); col += 1
        ws.cell(row=row_idx, column=col, value=r["motion"]); col += 1
        ws.cell(row=row_idx, column=col, value=r["argument_text"]); col += 1
        ws.cell(row=row_idx, column=col, value=r["overall_score"]); col += 1
        for cname in criteria_names:
            score, _ = r["criteria"].get(cname, (None, ""))
            ws.cell(row=row_idx, column=col, value=score); col += 1
        for cname in criteria_names:
            _, reasoning = r["criteria"].get(cname, (None, ""))
            ws.cell(row=row_idx, column=col, value=reasoning); col += 1
        ws.cell(row=row_idx, column=col, value="; ".join(r["strengths"])); col += 1
        ws.cell(row=row_idx, column=col, value="; ".join(r["weaknesses"])); col += 1
        ws.cell(row=row_idx, column=col, value="; ".join(r["suggestions"])); col += 1
        ws.cell(row=row_idx, column=col, value=r["elapsed_sec"]); col += 1
        col += 1  # Diem nguoi cham - de trong, dien tay sau
        col += 1  # Ghi chu - de trong
        ws.cell(row=row_idx, column=col, value=r["error"])

        for c in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=c).border = BORDER
            ws.cell(row=row_idx, column=c).alignment = WRAP
        ws.row_dimensions[row_idx].height = 60

    widths = [16, 26, 16, 34, 40, 16] + [10] * 5 + [34] * 5 + [30, 30, 30, 10, 18, 24, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Summary")
    scores = [r["overall_score"] for r in results if r["overall_score"] is not None]
    ws2["A1"] = "Batch Evaluation Summary"
    ws2["A1"].font = Font(name=FONT, size=13, bold=True)
    summary = [
        ("Total cases", len(results)),
        ("Succeeded", len(scores)),
        ("Failed", len(results) - len(scores)),
        ("Average overall_score", round(sum(scores) / len(scores), 2) if scores else "-"),
        ("Min overall_score", min(scores) if scores else "-"),
        ("Max overall_score", max(scores) if scores else "-"),
    ]
    for i, (k, v) in enumerate(summary, start=3):
        ws2.cell(row=i, column=1, value=k).font = Font(name=FONT, bold=True)
        ws2.cell(row=i, column=2, value=v)
    ws2.column_dimensions["A"].width = 26
    ws2.column_dimensions["B"].width = 16

    wb.save(path)
    print(f"\nDa xuat ket qua ra: {path}")


def main():
    settings = get_settings()
    client = get_llm_client(settings)
    cases = load_test_cases()
    print(f"Dang chay {len(cases)} test case voi provider = {settings.llm_provider} ...\n")
    results = run_all_cases(cases, settings, client)
    export_to_excel(results, OUTPUT_PATH)


if __name__ == "__main__":
    main()
