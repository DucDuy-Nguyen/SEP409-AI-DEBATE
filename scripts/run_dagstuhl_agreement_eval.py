"""
Chay bo dataset Dagstuhl-15512-ArgQuality Corpus (v2) qua AI Evaluator, roi tinh
Pearson r + p-value + MAE cho 6 CAP so sanh rieng biet (khong chi overall_score
nhu run_ibm_agreement_eval.py) — de biet CHINH XAC tieu chi nao AI dang khop
nguoi cham that, tieu chi nao dang lech nhieu nhat.

Nguon du lieu: file phang (KHONG dung thu muc XMI — XMI chi co 304/320 argument,
CSV phang co du 320):
    data/_dagstuhl_raw/extracted/dagstuhl-15512-argquality-corpus-v2/
        dagstuhl-15512-argquality-corpus-annotated.csv

Cach dung:
    python scripts/run_dagstuhl_agreement_eval.py --n 20 --runs 2

Yeu cau: da dien LLM key that trong .env. Da giai nen san file zip vao
data/_dagstuhl_raw/extracted/... (xem huong dan tai file trong lich su chat).

QUAN TRONG: giong run_ibm_agreement_eval.py, neu bi loi/ngat giua chung, script
VAN XUAT ra file Excel voi nhung bai da cham duoc, khong mat trang.
"""
import argparse
import csv
import json
import re
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.config import get_settings, Settings
from app.schemas import ArgumentEvaluationRequest, DebateSide, DebateStage
from app.services.evaluator import evaluate_argument
from app.services.llm_client import LLMClient, get_llm_client

RAW_CSV_PATH = (
    Path(__file__).parent.parent
    / "data" / "_dagstuhl_raw" / "extracted" / "dagstuhl-15512-argquality-corpus-v2"
    / "dagstuhl-15512-argquality-corpus-annotated.csv"
)
# File nguon dung dau xuong dong kieu Mac cu (chi \r, khong \n) — PHAI mo voi
# newline='' roi de csv.reader tu tach dong, KHONG duoc doc bang for-line-in-file
# thong thuong (se coi ca file la 1 dong duy nhat).
RAW_CSV_ENCODING = "cp1252"  # utf-8 loi tai 1 byte bi hong trong file goc; cp1252 doc duoc toan bo

PROCESSED_CACHE_PATH = Path(__file__).parent.parent / "data" / "dagstuhl_processed.csv"
OUTPUT_PATH = Path(__file__).parent.parent / "dagstuhl_agreement_results.xlsx"
DELAY_BETWEEN_CALLS_SEC = 1.5

# Cac cot BAT BUOC phai co trong CSV goc (SAU KHI chuan hoa ten cot: lowercase +
# strip + thay dau cach bang gach duoi). 6 cot dimension duoc chon tuong ung
# 1-1 voi thu tu tag trong XMI (da doi chieu thu cong — xem ghi chu ben duoi):
#   ArgumentationQuality -> "overall quality"   (khop AI overall_score)
#   Cogency              -> "cogency"           (khop AI Logic)
#   LocalSufficiency      -> "sufficiency"      (khop AI Evidence — LUU Y: cot
#                                                 CSV chi ten la "sufficiency",
#                                                 KHONG PHAI "local sufficiency"
#                                                 (ten do khong ton tai); phai
#                                                 phan biet voi cot "global
#                                                 sufficiency" la 1 dimension
#                                                 KHAC. Xac nhan bang cach doi
#                                                 chieu thu tu 15 tag trong file
#                                                 .xmi mau, giong het thu tu 15
#                                                 cot dimension trong CSV nay.)
#   GlobalRelevance       -> "global relevance"  (khop AI Relevance)
#   Arrangement           -> "arrangement"       (khop AI Structure)
#   Effectiveness         -> "effectiveness"     (khop AI Persuasiveness)
REQUIRED_COLUMNS = [
    "annotator", "argumentative", "argument", "#id", "issue",
    "overall_quality", "cogency", "sufficiency", "global_relevance", "arrangement", "effectiveness",
]

# dimension_key (dung noi bo + trong file cache) -> ten cot DA CHUAN HOA trong CSV goc
DIMENSION_COLUMNS = {
    "overall_quality": "overall_quality",
    "cogency": "cogency",
    "sufficiency": "sufficiency",
    "global_relevance": "global_relevance",
    "arrangement": "arrangement",
    "effectiveness": "effectiveness",
}

# 6 cap so sanh: (nhan hien thi, ten field AI, dimension_key nguoi tuong ung)
# ai_field = "overall_score" doc truc tiep tu result.overall_score (da o thang 0-10);
# cac ai_field con lai la ten 1 trong 5 CriterionScore.name (thang 1-5, se quy doi
# sang 0-10 rieng trong script nay CHI DE SO SANH — khong dung lai cong thuc nay
# de tinh overall_score that trong evaluator.py, cai do da doi sang holistic doc lap).
COMPARISON_PAIRS = [
    ("Overall (overall_score vs ArgumentationQuality)", "overall_score", "overall_quality"),
    ("Logic vs Cogency", "Logic", "cogency"),
    ("Evidence vs LocalSufficiency", "Evidence", "sufficiency"),
    ("Relevance vs GlobalRelevance", "Relevance", "global_relevance"),
    ("Structure vs Arrangement", "Structure", "arrangement"),
    ("Persuasiveness vs Effectiveness", "Persuasiveness", "effectiveness"),
]
AI_FIELDS = ["overall_score", "Logic", "Evidence", "Relevance", "Structure", "Persuasiveness"]

_SCORE_LABEL_RE = re.compile(r"^\s*(\d+(?:\.\d+)?)")

# Nhan KHONG PHAI so nhung VAN la gia tri hop le cua dataset (annotator co the
# cham 1 dimension cu the la "khong the danh gia" ke ca khi da xac nhan ca argument
# la argumentative='y') — phat hien qua khao sat thuc te toan bo gia tri trong 6
# cot dimension (chi co 5 gia tri khac nhau: "", "1 (Low)", "2 (Average)",
# "3 (High)", "Cannot judge"). Coi day la diem KHONG HOP LE, bo qua giong o rong,
# KHONG phai loi dinh dang can raise.
_NON_NUMERIC_VALID_LABELS = {"cannot judge"}


def _normalize_column_name(name: str) -> str:
    return name.strip().lower().replace(" ", "_")


def _parse_score_label(raw: str) -> float | None:
    """
    Chuyen nhan diem ve so. Dataset dung dinh dang "1 (Low)" / "2 (Average)" /
    "3 (High)", nhung ham nay TU NHAN DIEN: chi can lay so o DAU chuoi, nen cung
    xu ly duoc dinh dang khac (vi du chi co "2" hoac "2.0", khong co chu).
    Tra ve None neu chuoi rong (truong hop argumentative = 'n').
    """
    raw = (raw or "").strip()
    if not raw:
        return None
    if raw.strip().lower() in _NON_NUMERIC_VALID_LABELS:
        return None
    match = _SCORE_LABEL_RE.match(raw)
    if not match:
        raise ValueError(f"Khong nhan dien duoc dinh dang diem: {raw!r}")
    return float(match.group(1))


def human_score_1_3_to_10(mean_1_3: float) -> float:
    """Quy doi tuyen tinh thang 1-3 -> 0-10, cung y tuong voi human_score_to_10
    (thang 1-5) trong run_ibm_agreement_eval.py."""
    return round((mean_1_3 - 1) / 2 * 10, 2)


def _ai_criterion_1_5_to_10(score_1_5: float) -> float:
    """
    Quy doi diem 1 tieu chi con (1-5) sang thang 0-10 — CHI de so sanh cong bang
    voi thang 0-10 cua human trong script validate nay, KHONG lien quan gi den
    cach tinh overall_score that trong app/services/evaluator.py (da doi sang
    holistic doc lap, khong con dung cong thuc nay).
    """
    return round((score_1_5 - 1) / 4 * 10, 2)


def _build_column_index(header: list[str]) -> dict[str, int]:
    normalized = [_normalize_column_name(h) for h in header]
    index = {name: i for i, name in enumerate(normalized)}
    missing = [c for c in REQUIRED_COLUMNS if c not in index]
    if missing:
        raise RuntimeError(
            f"CSV Dagstuhl thieu cac cot bat buoc: {missing}.\n"
            f"Danh sach cot THAT da doc duoc (sau khi chuan hoa lowercase/strip/space->_): {normalized}"
        )
    return index


def build_processed_cache_if_needed(force: bool = False) -> None:
    """
    Doc file CSV goc (tab-separated, dau xuong dong kieu Mac \\r), gom 3 dong/
    annotator thanh 1 argument, tinh diem trung binh 6 dimension can dung, quy
    doi sang thang 0-10, roi luu ra data/dagstuhl_processed.csv (comma-separated
    binh thuong) de lan sau khong can parse lai CSV goc.
    """
    if PROCESSED_CACHE_PATH.exists() and not force:
        print(f"Da co san cache da xu ly: {PROCESSED_CACHE_PATH}, bo qua parse lai tu CSV goc.")
        return

    if not RAW_CSV_PATH.exists():
        raise FileNotFoundError(
            f"Khong tim thay file CSV goc: {RAW_CSV_PATH}\n"
            "Hay giai nen dagstuhl-15512-argquality-corpus-v2.zip vao dung duong dan nay truoc."
        )

    print(f"Dang parse CSV goc: {RAW_CSV_PATH} ...")
    dimension_keys = list(DIMENSION_COLUMNS.keys())

    # QUAN TRONG: file goc dung \r lam dau xuong dong (kieu Mac cu, khong phai
    # \n) — mo voi newline='' de csv module tu xu ly dung, KHONG dung
    # "for line in f" thong thuong (se doc ca file thanh 1 dong).
    with open(RAW_CSV_PATH, "r", newline="", encoding=RAW_CSV_ENCODING) as f:
        reader = csv.reader(f, delimiter="\t")
        header = next(reader)
        idx = _build_column_index(header)

        arguments: dict[str, dict] = {}
        for row in reader:
            if not row or len(row) <= max(idx.values()):
                continue  # bo qua dong rong/thieu cot o cuoi file (neu co)

            arg_id = row[idx["#id"]].strip()
            if not arg_id:
                continue

            entry = arguments.setdefault(arg_id, {
                "issue": row[idx["issue"]].strip(),
                "text": None,
                "n_valid_annotators": 0,
                "scores": {d: [] for d in dimension_keys},
            })

            if not entry["text"]:
                text = row[idx["argument"]].strip()
                if text:
                    entry["text"] = text

            argumentative = row[idx["argumentative"]].strip().lower()
            if argumentative != "y":
                # Dong nay khong hop le de tinh diem — BO QUA hoan toan (khong
                # phai coi la diem 0), dung nhu yeu cau.
                continue

            entry["n_valid_annotators"] += 1
            for dim_key, col_name in DIMENSION_COLUMNS.items():
                parsed = _parse_score_label(row[idx[col_name]])
                if parsed is not None:
                    entry["scores"][dim_key].append(parsed)

    rows_out = []
    skipped_no_valid_annotator = 0
    for arg_id, entry in arguments.items():
        if entry["n_valid_annotators"] == 0:
            skipped_no_valid_annotator += 1
            continue  # tat ca annotator deu cham "khong phai argument" -> bo qua

        row_out = {
            "id": arg_id,
            "issue": entry["issue"],
            "text": entry["text"] or "",
            "n_valid_annotators": entry["n_valid_annotators"],
        }
        for dim_key in dimension_keys:
            scores = entry["scores"][dim_key]
            if scores:
                mean_1_3 = sum(scores) / len(scores)
                row_out[f"{dim_key}_mean_1_3"] = round(mean_1_3, 3)
                row_out[f"{dim_key}_score_10"] = human_score_1_3_to_10(mean_1_3)
            else:
                row_out[f"{dim_key}_mean_1_3"] = ""
                row_out[f"{dim_key}_score_10"] = ""
        rows_out.append(row_out)

    fieldnames = ["id", "issue", "text", "n_valid_annotators"]
    for dim_key in dimension_keys:
        fieldnames += [f"{dim_key}_mean_1_3", f"{dim_key}_score_10"]

    PROCESSED_CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(PROCESSED_CACHE_PATH, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows_out)

    print(
        f"Da xu ly {len(rows_out)} argument hop le "
        f"(bo qua {skipped_no_valid_annotator} argument khong co annotator nao hop le). "
        f"Da luu cache: {PROCESSED_CACHE_PATH}"
    )


def load_dagstuhl_sample(n: int, seed: int = 42) -> pd.DataFrame:
    df = pd.read_csv(PROCESSED_CACHE_PATH)
    n = min(n, len(df))
    return df.sample(n=n, random_state=seed).reset_index(drop=True)


def _truncate(text: str | None, limit: int = 100) -> str:
    if not text:
        return ""
    if len(text) <= limit:
        return text
    return text[:limit].rstrip() + "..."


def _evaluate_one_row(row, settings, client, n_runs: int) -> dict:
    """Cham 1 argument N_RUNS lan doc lap, gom ket qua cho CA 6 field AI can
    theo doi (overall_score + 5 tieu chi con, da quy doi ve thang 0-10)."""
    entry = {
        "id": row["id"],
        "issue": row["issue"],
        "text": row["text"],
        "n_valid_annotators": int(row["n_valid_annotators"]),
        "human_scores_10": {},
    }
    for _, ai_field, dim_key in COMPARISON_PAIRS:
        val = row.get(f"{dim_key}_score_10")
        entry["human_scores_10"][dim_key] = float(val) if pd.notna(val) and val != "" else None

    runs_scores: dict[str, list[float]] = {field: [] for field in AI_FIELDS}
    overall_reasoning_list: list[str] = []

    req = ArgumentEvaluationRequest(
        motion=row["issue"],
        # Dataset khong co nhan Pro/Con dung enum DebateSide — MAC DINH dung
        # side=PRO cho TAT CA. O ArgumentContext.SOLO (khong co
        # opponent_argument_text), side chi anh huong framing cau chu trong
        # prompt (vi du "Phia dang dung: Ung ho (Pro)"), KHONG anh huong ban
        # chat cach cham Logic/Evidence/Structure/Persuasiveness/Relevance —
        # nen chon PRO co dinh khong lam sai lech ket qua so sanh.
        side=DebateSide.PRO,
        # Dataset la argument doc lap, khong co cau truc doi dap -> giong IBM
        # dataset -> opponent_argument_text=None -> ArgumentContext.SOLO ->
        # dung dung rubric Relevance da bo phan "phan bien doi phuong".
        stage=DebateStage.OPENING,
        argument_text=row["text"],
        opponent_argument_text=None,
    )

    for run_idx in range(1, n_runs + 1):
        try:
            result = evaluate_argument(req, settings=settings, llm_client=client)
            runs_scores["overall_score"].append(result.overall_score)
            for c in result.criteria:
                if c.name in runs_scores:
                    runs_scores[c.name].append(_ai_criterion_1_5_to_10(c.score))
            overall_reasoning_list.append(result.overall_reasoning)
            print(f"run{run_idx}=OK", end=" ")
        except Exception as e:
            print(f"run{run_idx}=LOI({type(e).__name__})", end=" ")
        if run_idx < n_runs:
            time.sleep(DELAY_BETWEEN_CALLS_SEC)

    entry["ai_avg"] = {}
    entry["ai_range"] = {}
    for field, scores in runs_scores.items():
        if scores:
            entry["ai_avg"][field] = round(sum(scores) / len(scores), 2)
            entry["ai_range"][field] = round(max(scores) - min(scores), 2)
        else:
            entry["ai_avg"][field] = None
            entry["ai_range"][field] = None

    entry["overall_reasoning"] = overall_reasoning_list[0] if overall_reasoning_list else None
    entry["error"] = None
    return entry


def run_batch(df: pd.DataFrame, settings, client, n_runs: int, rows: list) -> None:
    """Ghi truc tiep vao `rows` truyen vao (khong tao list moi roi return) —
    giong run_ibm_agreement_eval.py, de neu loi/Ctrl+C giua chung, ket qua da
    cham duoc khong bi mat."""
    for i, row in df.iterrows():
        print(f"[{i + 1}/{len(df)}] id={row['id']} ...", end=" ")
        try:
            entry = _evaluate_one_row(row, settings, client, n_runs)
            rows.append(entry)
            reasoning_preview = _truncate(entry.get("overall_reasoning"), limit=100)
            print(f"| reasoning: {reasoning_preview}")
        except Exception as e:
            print(f"LOI TOAN BO DONG: {e}")
            rows.append({
                "id": row.get("id"), "issue": row.get("issue"), "text": row.get("text"),
                "n_valid_annotators": row.get("n_valid_annotators"),
                "human_scores_10": {}, "ai_avg": {}, "ai_range": {},
                "overall_reasoning": None, "error": str(e),
            })


def compute_agreement_metrics(rows: list[dict]) -> dict[str, dict]:
    """Pearson r, p-value, MAE cho TUNG cap trong COMPARISON_PAIRS rieng biet."""
    from scipy import stats

    metrics = {}
    for label, ai_field, dim_key in COMPARISON_PAIRS:
        pairs = []
        for r in rows:
            human_val = r.get("human_scores_10", {}).get(dim_key)
            ai_val = r.get("ai_avg", {}).get(ai_field)
            if human_val is not None and ai_val is not None:
                pairs.append((human_val, ai_val))

        if len(pairs) < 2:
            metrics[label] = {"n": len(pairs), "pearson_r": None, "p_value": None, "mae": None}
            continue

        human_scores = [p[0] for p in pairs]
        ai_scores = [p[1] for p in pairs]
        pearson_r, p_value = stats.pearsonr(human_scores, ai_scores)
        mae = sum(abs(h - a) for h, a in pairs) / len(pairs)

        metrics[label] = {
            "n": len(pairs),
            "pearson_r": round(pearson_r, 3),
            "p_value": round(p_value, 4),
            "mae": round(mae, 2),
        }
    return metrics


FONT = "Arial"
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def _write_header(ws, headers: list[str]):
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.freeze_panes = "A2"


def export_to_excel(rows: list[dict], metrics: dict, path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Agreement Results"

    pair_headers = []
    for label, ai_field, dim_key in COMPARISON_PAIRS:
        pair_headers += [f"{label} - AI avg", f"{label} - AI range", f"{label} - Human", f"{label} - Diff"]

    headers = (
        ["Arg ID", "Issue (motion)", "Text (truncated)", "# Valid annotators"]
        + pair_headers
        + ["Overall Reasoning (run 1)", "Error"]
    )
    _write_header(ws, headers)

    for row_idx, r in enumerate(rows, start=2):
        col = 1
        ws.cell(row=row_idx, column=col, value=r.get("id")); col += 1
        ws.cell(row=row_idx, column=col, value=r.get("issue")); col += 1
        text = r.get("text") or ""
        ws.cell(row=row_idx, column=col, value=text[:300]); col += 1
        ws.cell(row=row_idx, column=col, value=r.get("n_valid_annotators")); col += 1

        for label, ai_field, dim_key in COMPARISON_PAIRS:
            ai_avg = (r.get("ai_avg") or {}).get(ai_field)
            ai_range = (r.get("ai_range") or {}).get(ai_field)
            human_val = (r.get("human_scores_10") or {}).get(dim_key)
            diff = round(ai_avg - human_val, 2) if ai_avg is not None and human_val is not None else None
            ws.cell(row=row_idx, column=col, value=ai_avg); col += 1
            ws.cell(row=row_idx, column=col, value=ai_range); col += 1
            ws.cell(row=row_idx, column=col, value=human_val); col += 1
            ws.cell(row=row_idx, column=col, value=diff); col += 1

        ws.cell(row=row_idx, column=col, value=r.get("overall_reasoning")); col += 1
        ws.cell(row=row_idx, column=col, value=r.get("error"))

        for c in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=c).border = BORDER
            ws.cell(row=row_idx, column=c).alignment = WRAP
        ws.row_dimensions[row_idx].height = 50

    widths = [10, 34, 44, 14] + [12, 10, 10, 10] * len(COMPARISON_PAIRS) + [50, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("Agreement Metrics")
    ws2["A1"] = "Agreement Rate — AI vs Human (Dagstuhl-15512-ArgQuality corpus, 6 cap so sanh)"
    ws2["A1"].font = Font(name=FONT, size=13, bold=True)

    metrics_headers = ["Comparison", "N", "Pearson r", "p-value", "MAE (0-10)"]
    for col, h in enumerate(metrics_headers, start=1):
        cell = ws2.cell(row=3, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER

    for i, (label, ai_field, dim_key) in enumerate(COMPARISON_PAIRS, start=4):
        m = metrics.get(label, {})
        ws2.cell(row=i, column=1, value=label)
        ws2.cell(row=i, column=2, value=m.get("n"))
        ws2.cell(row=i, column=3, value=m.get("pearson_r"))
        ws2.cell(row=i, column=4, value=m.get("p_value"))
        ws2.cell(row=i, column=5, value=m.get("mae"))
        for c in range(1, 6):
            ws2.cell(row=i, column=c).border = BORDER

    note_row = 4 + len(COMPARISON_PAIRS) + 1
    notes = [
        ("Cach doc ket qua:", ""),
        ("r > 0.5", "Tuong quan kha, AI va nguoi co xu huong danh gia giong nhau cho tieu chi nay"),
        ("r 0.3 - 0.5", "Tuong quan yeu, can xem lai mo ta tieu chi/prompt tuong ung"),
        ("r < 0.3", "Gan nhu khong tuong quan, AI dang danh gia khac han nguoi o tieu chi nay"),
        ("p-value < 0.05", "Ket qua co y nghia thong ke (khong phai ngau nhien)"),
    ]
    for i, (k, v) in enumerate(notes, start=note_row):
        ws2.cell(row=i, column=1, value=k).font = Font(name=FONT, bold=(v == ""))
        ws2.cell(row=i, column=2, value=v)

    ws2.column_dimensions["A"].width = 46
    ws2.column_dimensions["B"].width = 10
    ws2.column_dimensions["C"].width = 12
    ws2.column_dimensions["D"].width = 12
    ws2.column_dimensions["E"].width = 12

    wb.save(path)
    print(f"\nDa xuat ket qua ra: {path}")


def main():
    # overall_reasoning tu LLM la tieng Viet co dau — console Windows mac dinh
    # dung codepage (vi du cp1252) khong encode duoc, se crash print() giua
    # chung va lam mat/hong du lieu dong dang cham (xem bai hoc tu
    # run_ibm_agreement_eval.py). Ep stdout sang UTF-8 de tranh loi nay.
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

    parser = argparse.ArgumentParser()
    parser.add_argument("--n", type=int, default=20, help="So luong argument lay mau de test")
    parser.add_argument("--seed", type=int, default=42, help="Random seed de lay mau on dinh, lap lai duoc")
    parser.add_argument("--runs", type=int, default=2, help="So lan cham MOI argument (de do dao dong)")
    args = parser.parse_args()

    build_processed_cache_if_needed()

    settings = get_settings()
    client = get_llm_client(settings)

    df = load_dagstuhl_sample(n=args.n, seed=args.seed)
    print(
        f"\nDang cham {len(df)} argument tu Dagstuhl-15512-ArgQuality x {args.runs} lan/bai "
        f"voi provider = {settings.llm_provider} ...\n"
    )

    rows: list[dict] = []
    try:
        run_batch(df, settings, client, n_runs=args.runs, rows=rows)
    except KeyboardInterrupt:
        print("\n\nBi ngat boi nguoi dung (Ctrl+C) — se xuat ket qua da cham duoc tinh den day.")
    except Exception as e:
        print(f"\n\nDung giua chung do loi khong luong truoc: {e}")
        print("Van se xuat ket qua da cham duoc tinh den day.")
    finally:
        if rows:
            metrics = compute_agreement_metrics(rows)
            print("\n=== KET QUA AGREEMENT (6 CAP SO SANH, tren so argument da cham xong) ===")
            print(json.dumps(metrics, indent=2, ensure_ascii=False))
            export_to_excel(rows, metrics, OUTPUT_PATH)
        else:
            print("\nChua cham duoc argument nao, khong co gi de xuat.")


if __name__ == "__main__":
    main()
